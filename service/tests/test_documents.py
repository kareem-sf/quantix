import io
import time

import docx
import openpyxl
import pytest


def make_pdf(pages: list[list[str]]) -> bytes:
    """A small, valid PDF with one Helvetica text line per entry; an empty page stands in for a scan."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for lines in pages:
        stream = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({line}) '" for line in lines) + " ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content = len(objects)
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content} 0 R >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


def make_xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BOQ"
    sheet.append(["Item", "Description", "Unit", "Qty"])
    sheet.append(["3.1", "Excavation to reduce levels", "m3", 1240.0])
    sheet.append(["4.2", "خرسانة مُسلّحة للبلاطة", "م3", 312.4])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Instructions to Tenderers")
    document.add_paragraph("The bid bond shall be one percent of the tender price.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Validity"
    table.rows[0].cells[1].text = "120 days"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


PDF = make_pdf([["Conditions of Contract", "4.2 Tender security of one percent"], []])


@pytest.fixture
def tender(client):
    return client.post("/tenders", json={"name": "Synthetic school"}).json()["id"]


def upload(client, tender_id, files: dict[str, bytes]):
    return client.post(
        f"/tenders/{tender_id}/documents",
        files=[("files", (name, content, "application/octet-stream")) for name, content in files.items()],
    )


def read_all(client, tender_id):
    for _ in range(100):
        documents = client.get(f"/tenders/{tender_id}/documents").json()
        if all(d["status"] not in ("waiting", "reading") for d in documents):
            return {d["path"]: d for d in documents}
        time.sleep(0.05)
    raise AssertionError("The documents were not read in time.")


def test_a_package_is_stored_and_read(client, tender, tmp_path):
    response = upload(
        client,
        tender,
        {
            "Package/Conditions.pdf": PDF,
            "Package/BOQ/Bill.xlsx": make_xlsx(),
            "Package/ITT.docx": make_docx(),
            "Package/Drawings/A-101.dwg": b"AC1027 not really a drawing",
        },
    )
    assert response.json() == {"added": 4, "unchanged": 0}
    documents = read_all(client, tender)

    pdf = documents["Package/Conditions.pdf"]
    assert (pdf["status"], pdf["page_count"], pdf["name"]) == ("read", 2, "Conditions.pdf")
    assert pdf["note"] == "1 of 2 pages are scans without text. The office reads them from the page image."
    first = client.get(f"/documents/{pdf['id']}/pages/1").json()
    assert "Tender security of one percent" in first["text"] and first["has_text"] is True
    assert client.get(f"/documents/{pdf['id']}/pages/2").json()["has_text"] is False

    sheet = client.get(f"/documents/{documents['Package/BOQ/Bill.xlsx']['id']}/pages/1").json()["text"]
    assert "Sheet: BOQ" in sheet and "A2=3.1 | B2=Excavation to reduce levels | C2=m3 | D2=1240" in sheet

    word = client.get(f"/documents/{documents['Package/ITT.docx']['id']}/pages/1").json()["text"]
    assert "one percent of the tender price" in word and "Validity | 120 days" in word

    dwg = documents["Package/Drawings/A-101.dwg"]
    assert dwg["status"] == "unreadable" and dwg["note"].startswith("CAD drawings can't be read yet")
    kept = (tmp_path / "tenders" / tender / "files").iterdir()
    assert sorted(f.suffix for f in kept) == [".docx", ".dwg", ".pdf", ".xlsx"]


def test_search_finds_pages_in_english_and_arabic(client, tender):
    upload(client, tender, {"Conditions.pdf": PDF, "Bill.xlsx": make_xlsx()})
    read_all(client, tender)

    hits = client.get(f"/tenders/{tender}/search", params={"q": "tender security"}).json()
    assert [(h["name"], h["page"]) for h in hits] == [("Conditions.pdf", 1)]
    assert "4.2 [Tender] [security] of one percent" == hits[0]["snippet"]

    arabic = client.get(f"/tenders/{tender}/search", params={"q": "خرسانة مسلحة"}).json()
    assert [(h["name"], h["page"]) for h in arabic] == [("Bill.xlsx", 1)]
    assert "مُسلّحة" in arabic[0]["snippet"]  # the document's own spelling, marks included
    assert client.get(f"/tenders/{tender}/search", params={"q": "nothing like this"}).json() == []


def test_the_same_file_is_ignored_and_a_changed_one_replaces_it(client, tender):
    upload(client, tender, {"Conditions.pdf": PDF})
    assert upload(client, tender, {"Conditions.pdf": PDF}).json() == {"added": 0, "unchanged": 1}

    upload(client, tender, {"Conditions.pdf": make_pdf([["Revised conditions"]])})
    read_all(client, tender)
    documents = client.get(f"/tenders/{tender}/documents").json()
    assert sorted(d["status"] for d in documents) == ["read", "replaced"]
    current = next(d for d in documents if d["status"] == "read")
    assert "Revised conditions" in client.get(f"/documents/{current['id']}/pages/1").json()["text"]


def test_page_images_and_originals(client, tender):
    upload(client, tender, {"Conditions.pdf": PDF, "ITT.docx": make_docx()})
    documents = read_all(client, tender)

    image = client.get(f"/documents/{documents['Conditions.pdf']['id']}/pages/1/image")
    assert image.headers["content-type"] == "image/png" and image.content.startswith(b"\x89PNG")
    assert client.get(f"/documents/{documents['ITT.docx']['id']}/pages/1/image").status_code == 404
    assert client.get(f"/documents/{documents['Conditions.pdf']['id']}/file").content == PDF


def test_rejects_paths_that_leave_the_package(client, tender):
    assert upload(client, tender, {"../outside.pdf": PDF}).status_code == 400
