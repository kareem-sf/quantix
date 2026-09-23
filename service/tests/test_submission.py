import hashlib
import io
from decimal import Decimal

import docx
import openpyxl
import pytest
from pydantic_ai.messages import ModelResponse, TextPart, ToolCallPart, ToolReturnPart
from pydantic_ai.models.function import FunctionModel
from test_documents import make_pdf, read_all, upload
from test_office import wait_for

from quantix import settings
from quantix.boq import records as boq
from quantix.documents import library
from quantix.documents.models import Document
from quantix.estimate import records as estimate
from quantix.office import records as office
from quantix.submission import export
from quantix.submission import records as submission

ITT = make_pdf(
    [
        [
            "Instructions to Tenderers",
            "7.1 The priced bill of quantities shall be submitted in the format provided.",
            "7.3 A bid bond of 1% of the tender price shall accompany the tender.",
            "7.6 The tenderer shall submit a method statement for concrete works.",
        ]
    ]
)


def client_bill() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BOQ"
    sheet.append(["Item", "Description", "Unit", "Qty", "Rate", "Amount"])
    sheet.append(["3.1", "Excavation", "m3", 1240, None, None])
    sheet.append(["4.3", "Slab reinforcement", "t", 28.1, None, "=D3*E3"])
    sheet.append(["6.3", "Waterproofing", "m2", 980, None, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def tender(client):
    tender_id = client.post("/tenders", json={"name": "Synthetic school"}).json()["id"]
    upload(client, tender_id, {"Bill.xlsx": client_bill(), "ITT.pdf": ITT})
    docs = read_all(client, tender_id)
    with client.app.state.sessions() as session:
        office.hire(session, tender_id, "Rania Farouk", "Tender Manager", {}, is_manager=True)
        layla = office.hire(session, tender_id, "Layla Nasser", "Tender Coordinator", {})
        rows = [("3.1", "Excavation", "m3", "1240"), ("4.3", "Slab reinforcement", "t", "28.1")]
        rows.append(("6.3", "Waterproofing", "m2", "980"))
        lines = [
            boq.ItemIn(
                item=i,
                description=d,
                unit=u,
                quantity=Decimal(q),
                document_id=docs["Bill.xlsx"]["id"],
                page=1,
                quote=f"A{n}={i} | B{n}={d} | C{n}={u} | D{n}={q}",
            )
            for n, (i, d, u, q) in enumerate(rows, start=2)
        ]
        boq.propose_items(session, tender_id, layla, lines, False)
        boq.approve_all_items(session, tender_id)
        note = "Own rate from outputs and current prices."
        for item, rate in (("3.1", "18.50"), ("4.3", "3488.00")):
            estimate.propose_rate(
                session, tender_id, layla.id, item, "estimate", note, unit_rate=Decimal(rate), status="approved"
            )
        zero = Decimal(0)
        estimate.propose_markups(
            session, tender_id, layla.id, Decimal("0.10"), zero, zero, zero, "Site costs", status="approved"
        )
        session.commit()
        return tender_id, layla.id, docs["Bill.xlsx"]["id"], docs["ITT.pdf"]["id"]


def requirement(title, quote, itt, section="Commercial"):
    return submission.RequirementIn(section=section, title=title, document_id=itt, page=1, quote=quote)


def checklist(client, tender):
    tender_id, layla, _, itt = tender
    with client.app.state.sessions() as session:
        report = submission.add_requirements(
            session,
            tender_id,
            layla,
            [
                requirement("Bid bond, 1% of the tender price", "7.3 A bid bond of 1%", itt),
                requirement(
                    "Method statement for concrete works", "7.6 The tenderer shall submit a method", itt, "Technical"
                ),
                requirement("Site visit certificate", "7.9 A site visit certificate", itt),
            ],
        )
        session.commit()
    return report


def test_the_checklist_keeps_only_requirements_the_tender_states(client, tender):
    report = checklist(client, tender)
    assert report.splitlines()[0] == "2 requirements added to the checklist."
    assert report.splitlines()[1].startswith("Site visit certificate: “7.9 A site visit certificate” is not on ITT.pdf")
    rows = client.get(f"/tenders/{tender[0]}/submission").json()["requirements"]
    assert [(r["section"], r["title"], r["state"], r["document_name"], r["page"]) for r in rows] == [
        ("Commercial", "Bid bond, 1% of the tender price", "missing", "ITT.pdf", 1),
        ("Technical", "Method statement for concrete works", "missing", "ITT.pdf", 1),
    ]


def test_drafts_wait_for_review_and_the_engineer_marks_what_they_provide(client, tender):
    tender_id, layla, _, _ = tender
    checklist(client, tender)
    with client.app.state.sessions() as session:
        method = submission.find_requirement(session, tender_id, "method statement for concrete works")
        submission.draft(session, method, layla, "Method statement", "First try.")
        second = submission.draft(session, method, layla, "Method statement", "Pour sequence.\n\nCuring for 7 days.")
        session.commit()
    rows = {r["title"]: r for r in client.get(f"/tenders/{tender_id}/submission").json()["requirements"]}
    method, bond = rows["Method statement for concrete works"], rows["Bid bond, 1% of the tender price"]
    assert (method["state"], method["draft"]["id"], method["draft"]["body"]) == (
        "review",
        second.id,
        "Pour sequence.\n\nCuring for 7 days.",
    )
    assert client.get(f"/tenders/{tender_id}/gates").json()["submission"] == 1

    client.post(f"/drafts/{second.id}/decision", json={"approve": False, "reason": "Add the pour sizes."})
    chat = client.get(f"/tenders/{tender_id}/messages", params={"channel": layla}).json()
    assert chat[-1]["text"] == "I sent back the draft “Method statement”: Add the pour sizes."
    with client.app.state.sessions() as session:
        third = submission.draft(
            session, session.get(submission.Requirement, method["id"]), layla, "Method statement", "v3"
        )
        session.commit()
    client.post(f"/drafts/{third.id}/decision", json={"approve": True})
    client.post(f"/requirements/{bond['id']}/ready", json={"ready": True, "note": "The bank issues it on Monday."})
    rows = client.get(f"/tenders/{tender_id}/submission").json()["requirements"]
    assert [r["state"] for r in rows] == ["ready", "ready"]
    assert client.get(f"/tenders/{tender_id}/gates").json()["submission"] == 0

    client.post(f"/requirements/{bond['id']}/ready", json={"ready": False})
    signed = client.post(f"/requirements/{bond['id']}/file", files={"file": ("Bond.pdf", b"%PDF-1.4 signed")})
    assert signed.status_code == 200
    assert client.get(f"/tenders/{tender_id}/submission").json()["requirements"][0]["file_name"] == "Bond.pdf"


def test_pricing_columns_come_from_the_client_header(client, tender):
    tender_id, layla, bill, itt = tender
    with client.app.state.sessions() as session:
        with pytest.raises(ValueError, match="no cell in column G"):
            submission.set_pricing_columns(session, tender_id, layla, bill, 1, "E", "G", "E1=Rate | F1=Amount")
        with pytest.raises(ValueError, match="client's BOQ workbook"):
            submission.set_pricing_columns(session, tender_id, layla, itt, 1, "E", "F", "7.1 The priced bill")
        submission.set_pricing_columns(session, tender_id, layla, bill, 1, "e", "f", "E1=Rate | F1=Amount")
        session.commit()
    assert client.get(f"/tenders/{tender_id}/submission").json()["columns"] == [
        {
            "document_id": bill,
            "document_name": "Bill.xlsx",
            "sheet": 1,
            "rate_column": "E",
            "amount_column": "F",
            "proposed_by": layla,
        }
    ]


def test_the_package_is_built_in_the_client_format_with_markups_in_the_rates(client, tender, tmp_path):
    tender_id, layla, bill, _ = tender
    checklist(client, tender)
    with client.app.state.sessions() as session:
        submission.set_pricing_columns(session, tender_id, layla, bill, 1, "E", "F", "E1=Rate | F1=Amount")
        method = submission.find_requirement(session, tender_id, "Method statement for concrete works")
        submission.draft(session, method, layla, "Method statement", "Pour sequence.\n\nCuring.", "office_approved")
        stored = library.stored_file(tmp_path, session.get(Document, bill))
        session.commit()
    before = hashlib.sha256(stored.read_bytes()).hexdigest()

    built = client.post(f"/tenders/{tender_id}/export", json={"spread_markups": True}).json()
    # Net 1240 × 18.50 + 28.1 × 3488.00 = 22,940.00 + 98,012.80 = 120,952.80; preliminaries 10% = 12,095.28.
    # The total 133,048.08 ÷ 120,952.80 = 1.1, so the rates become 20.35 and 3,836.80.
    assert (Decimal(built["factor"]), built["summary_total"], built["priced_total"]) == (
        Decimal("1.1"),
        "133048.08",
        "133048.08",
    )
    assert built["files"] == ["Priced Bill.xlsx", "Method statement.docx", "Checklist.xlsx"]
    assert built["not_ready"] == ["BOQ item 6.3 is not priced", "Bid bond, 1% of the tender price"]

    folder = tmp_path / "exports" / built["folder"]
    assert built["folder"].startswith("Synthetic school ")
    sheet = openpyxl.load_workbook(folder / "Priced Bill.xlsx").active
    assert [[sheet[f"{c}{r}"].value for c in "EF"] for r in (2, 3, 4)] == [
        [20.35, 25234],
        [3836.8, "=D3*E3"],  # the client's own formula is kept
        [None, None],
    ]
    assert hashlib.sha256(stored.read_bytes()).hexdigest() == before  # the supplied file is unchanged
    paragraphs = [p.text for p in docx.Document(folder / "Method statement.docx").paragraphs]
    assert paragraphs == ["Method statement", "Pour sequence.", "Curing."]
    rows = list(openpyxl.load_workbook(folder / "Checklist.xlsx").active.values)
    assert rows[1:] == [
        ("Commercial", "Bid bond, 1% of the tender price", "ITT.pdf, page 1", "Missing", None),
        (
            "Technical",
            "Method statement for concrete works",
            "ITT.pdf, page 1",
            "Ready · approved by the office, not reviewed",
            "Method statement.docx",
        ),
    ]

    with client.app.state.sessions() as session:
        rates, factor = export.submitted_rates(session, tender_id, spread=False)
    assert (factor, sorted(rates.values())) == (Decimal(1), [Decimal("18.50"), Decimal("3488.00")])
    assert client.post("/exports/..%5C..%5Cwindows/open").status_code == 404


def test_staff_build_the_checklist_through_their_tools(client, tender, tmp_path):
    tender_id, layla, bill, itt = tender
    settings.save(tmp_path, office_ai={"connection_id": "scripted", "model": "brain"})
    replies: list[str] = []
    steps = [
        ToolCallPart(
            "add_requirements",
            {
                "requirements": [
                    {
                        "section": "Technical",
                        "title": "Method statement for concrete works",
                        "document_id": itt,
                        "page": 1,
                        "quote": "7.6 The tenderer shall submit a method statement for concrete works.",
                    }
                ]
            },
        ),
        ToolCallPart(
            "draft_document",
            {"requirement": "Method statement for concrete works", "title": "Method statement", "text": "Pour."},
        ),
        ToolCallPart(
            "set_pricing_columns",
            {
                "document_id": bill,
                "sheet": 1,
                "rate_column": "E",
                "amount_column": "F",
                "header_quote": "E1=Rate | F1=Amount",
            },
        ),
        ToolCallPart("list_requirements", {}),
    ]

    def brain(messages, info):
        if "You are Layla Nasser" not in info.instructions:
            return ModelResponse(parts=[TextPart("Done.")])
        done = [str(p.content) for m in messages for p in m.parts if isinstance(p, ToolReturnPart)]
        replies[:] = done
        return (
            ModelResponse(parts=[steps[len(done)]])
            if len(done) < len(steps)
            else ModelResponse(parts=[TextPart("Done.")])
        )

    client.app.state.office.model = lambda: FunctionModel(brain)
    with client.app.state.sessions() as session:
        office.post(session, tender_id, "engineer", layla, "Layla, start the submission checklist.")
        session.commit()
    client.app.state.office.engineer_spoke(tender_id)
    wait_for(lambda o: len(replies) == len(steps), client, tender_id)
    assert replies == [
        "1 requirements added to the checklist.",
        "The draft is waiting for the engineer.",
        "Rates will go in column E and amounts in column F.",
        "- Technical · Method statement for concrete works: review",
    ]
