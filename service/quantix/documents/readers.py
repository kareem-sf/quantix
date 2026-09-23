"""Read a tender file into numbered pages of text. PDF pages keep their page numbers; a spreadsheet gives one page
per sheet; a Word document gives pages of about 3,000 characters split at paragraphs."""

import ctypes
import threading
from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from pathlib import Path

import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_raw

from quantix.documents.arabic import Char, has_arabic, lines_from_boxes

# PDFium is not thread-safe: every use of it in the service goes through this lock.
PDFIUM = threading.Lock()

KINDS = {
    ".pdf": "pdf",
    ".xlsx": "spreadsheet",
    ".xlsm": "spreadsheet",
    ".docx": "word",
    ".doc": "old_word",
    ".xls": "old_spreadsheet",
    ".dwg": "cad",
    ".dxf": "cad",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tif": "image",
    ".tiff": "image",
}
UNREADABLE = {
    "old_word": "Old Word files (.doc) can't be read. Save it as .docx or PDF and add that copy.",
    "old_spreadsheet": "Old Excel files (.xls) can't be read. Save it as .xlsx and add that copy.",
    "cad": "CAD drawings can't be read yet. Add the PDF of this drawing instead.",
    "other": "This type of file isn't read.",
}
SCAN_CHARACTERS = 20  # a PDF page with fewer visible characters than this is treated as a scan


@dataclass
class PageText:
    number: int
    text: str
    has_text: bool = True
    width: float | None = None  # PDF points
    height: float | None = None


class Unreadable(Exception):
    """The file can't be read; the message says why, in plain words."""


def kind_of(path: str) -> str:
    return KINDS.get(Path(path).suffix.lower(), "other")


def read_file(path: Path, kind: str) -> list[PageText]:
    if kind in UNREADABLE:
        raise Unreadable(UNREADABLE[kind])
    if kind == "pdf":
        return _pdf(path)
    if kind == "spreadsheet":
        return _spreadsheet(path)
    if kind == "word":
        return _word(path)
    return [PageText(1, "", has_text=False)]  # an image: the office reads it by looking at it


def _pdf(path: Path) -> list[PageText]:
    with PDFIUM:
        try:
            document = pdfium.PdfDocument(path)
        except pdfium.PdfiumError as error:
            if "password" in str(error).lower():
                raise Unreadable("This PDF is protected with a password.") from error
            raise Unreadable("This PDF is damaged and can't be opened.") from error
        try:
            pages = []
            for index in range(len(document)):
                page = document[index]
                text = _pdf_page_text(page.get_textpage())
                visible = sum(not c.isspace() for c in text)
                width, height = page.get_size()
                pages.append(PageText(index + 1, text, visible >= SCAN_CHARACTERS, width, height))
            return pages
        finally:
            document.close()


def _pdf_page_text(textpage: pdfium.PdfTextPage) -> str:
    plain = textpage.get_text_bounded()
    if not has_arabic(plain):
        return plain.replace("\r\n", "\n").strip()
    chars = []
    for index in range(textpage.count_chars()):
        code = pdfium_raw.FPDFText_GetUnicode(textpage.raw, index)
        if not code:
            continue
        left, bottom, right, top = textpage.get_charbox(index, loose=True)  # font boxes: one height per line
        chars.append(Char(chr(code), left, bottom, right, top))
    return "\n".join(lines_from_boxes(chars))


def _spreadsheet(path: Path) -> list[PageText]:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise Unreadable("This workbook is damaged or protected and can't be opened.") from error
    pages = []
    try:
        for number, sheet in enumerate(workbook.worksheets, start=1):
            lines = [f"Sheet: {sheet.title}"]
            for row in sheet.iter_rows():
                cells = [f"{c.coordinate}={_cell(c.value)}" for c in row if c.value not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
            pages.append(PageText(number, "\n".join(lines)))
    finally:
        workbook.close()
    return pages


def _cell(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\n", " ").strip()


def _word(path: Path) -> list[PageText]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(BytesIO(path.read_bytes()))
    except Exception as error:
        raise Unreadable("This Word document is damaged or protected and can't be opened.") from error
    blocks: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph) and block.text.strip():
            blocks.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
    pages: list[PageText] = []
    current: list[str] = []
    for block in blocks:
        if current and sum(len(b) for b in current) + len(block) > 3000:
            pages.append(PageText(len(pages) + 1, "\n".join(current)))
            current = []
        current.append(block)
    if current or not pages:
        pages.append(PageText(len(pages) + 1, "\n".join(current)))
    return pages


def find_text(path: Path, number: int, text: str, limit: int = 20) -> list[tuple[float, float, float, float]]:
    """Where `text` is printed on a PDF page: (left, top, right, bottom) in points from the top left."""
    with PDFIUM:
        document = pdfium.PdfDocument(path)
        try:
            page = document[number - 1]
            height = page.get_height()
            textpage = page.get_textpage()
            searcher = textpage.search(text)
            boxes = []
            while len(boxes) < limit and (found := searcher.get_next()):
                index, count = found
                chars = [textpage.get_charbox(i) for i in range(index, index + count)]
                left, bottom = min(c[0] for c in chars), min(c[1] for c in chars)
                right, top = max(c[2] for c in chars), max(c[3] for c in chars)
                boxes.append((left, height - top, right, height - bottom))
            return boxes
        finally:
            document.close()


@lru_cache(maxsize=32)
def vector_points(path: Path, number: int, limit: int = 50_000) -> tuple[tuple[float, float], ...]:
    """The end and corner points of the lines drawn on a PDF page, in points from the top left. Takeoff snaps to
    them so measurements land exactly on the drawing. A scan has none."""
    found: set[tuple[float, float]] = set()
    with PDFIUM:
        document = pdfium.PdfDocument(path)
        try:
            page = document[number - 1]
            height = page.get_height()
            x, y = ctypes.c_float(), ctypes.c_float()
            for obj in page.get_objects(filter=[pdfium_raw.FPDF_PAGEOBJ_PATH], max_depth=5):
                matrix = obj.get_matrix()
                for index in range(pdfium_raw.FPDFPath_CountSegments(obj.raw)):
                    segment = pdfium_raw.FPDFPath_GetPathSegment(obj.raw, index)
                    pdfium_raw.FPDFPathSegment_GetPoint(segment, ctypes.byref(x), ctypes.byref(y))
                    px, py = matrix.on_point(x.value, y.value)
                    found.add((round(px, 4), round(height - py, 4)))
                    if len(found) >= limit:
                        return tuple(found)
        finally:
            document.close()
    return tuple(found)


def render_page(path: Path, number: int, width: int = 1400) -> bytes:
    """A PNG of one PDF page, about `width` pixels wide."""
    with PDFIUM:
        document = pdfium.PdfDocument(path)
        try:
            page = document[number - 1]
            image = page.render(scale=width / page.get_width()).to_pil()
        finally:
            document.close()
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
