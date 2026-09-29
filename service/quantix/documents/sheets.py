"""A worksheet as the Documents screen shows it: each cell's text as Excel displays it, with the sheet's merges,
column widths, row heights, fonts, fills, borders and alignment, so a workbook looks as it does in Excel. Formulas
show the value Excel saved with the file."""

import colorsys
import datetime as dt
import re
from functools import lru_cache
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from pydantic import BaseModel, ConfigDict

from quantix.documents.readers import Unreadable

MAX_ROWS = 5000  # rows past these are left to the original file
MAX_COLUMNS = 200
DEFAULT_WIDTH = 8.43  # Excel's default column width, in characters
DEFAULT_HEIGHT = 15.0  # and row height, in points
# Excel's theme colours in the order cells refer to them: light 1, dark 1, light 2, dark 2, accents 1-6, links
THEME = "lt1 dk1 lt2 dk2 accent1 accent2 accent3 accent4 accent5 accent6 hlink folHlink".split()
OFFICE_THEME = "FFFFFF 000000 E7E6E6 44546A 4472C4 ED7D31 A5A5A5 FFC000 5B9BD5 70AD47 0563C1 954F72".split()
BORDERS = {
    "hair": "1px solid",
    "thin": "1px solid",
    "dotted": "1px dotted",
    "dashed": "1px dashed",
    "dashDot": "1px dashed",
    "dashDotDot": "1px dashed",
    "medium": "2px solid",
    "mediumDashed": "2px dashed",
    "mediumDashDot": "2px dashed",
    "mediumDashDotDot": "2px dashed",
    "slantDashDot": "2px dashed",
    "thick": "3px solid",
    "double": "3px double",
}
ALIGN = {"left": "left", "center": "center", "centerContinuous": "center", "right": "right", "justify": "justify"}
VALIGN = {"top": "top", "center": "middle", "bottom": "bottom", "justify": "middle", "distributed": "middle"}


class Style(BaseModel):
    """How a cell looks. Colours are "#rrggbb"; a border side is CSS, e.g. "1px solid #000000"."""

    model_config = ConfigDict(frozen=True)  # cells that look alike share one

    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    font: str | None = None
    size: float | None = None  # points
    color: str | None = None
    fill: str | None = None
    align: str | None = None  # left | center | right | justify | end (a number, whichever way the sheet reads)
    valign: str | None = None  # top | middle | bottom
    wrap: bool = False
    indent: int = 0
    top: str | None = None
    right: str | None = None
    bottom: str | None = None
    left: str | None = None


class Cell(BaseModel):
    column: int  # its place among the columns shown
    text: str = ""
    style: int = 0  # into SheetView.styles
    rows: int = 1  # a merged cell: the rows and columns shown that it covers
    columns: int = 1


class Row(BaseModel):
    number: int  # Excel's row number
    height: float  # pixels
    hidden: bool = False  # hidden in the sheet, shown because hidden rows were asked for
    cells: list[Cell]  # the cells that start in this row, left to right; merged areas leave gaps


class Column(BaseModel):
    letter: str
    width: float  # pixels
    hidden: bool = False


class SheetTab(BaseModel):
    name: str
    hidden: bool = False


class SheetView(BaseModel):
    sheets: list[SheetTab]  # every sheet of the workbook, in order: sheet n is page n
    number: int
    right_to_left: bool = False
    gridlines: bool = True
    frozen_rows: int = 0  # rows and columns kept in view while scrolling
    frozen_columns: int = 0
    columns: list[Column]
    rows: list[Row]
    styles: list[Style]
    more_rows: int = 0  # rows past the last one shown
    more_columns: int = 0
    hidden_rows: int = 0  # rows and columns the sheet hides, shown only when asked for
    hidden_columns: int = 0


@lru_cache(maxsize=1)
def _workbook(path: Path):
    try:
        return load_workbook(path, data_only=True, keep_links=False)
    except Exception as error:
        raise Unreadable("This workbook is damaged or protected and can't be opened.") from error


@lru_cache(maxsize=16)  # stored files never change, so a sheet is worked out once
def view(path: Path, number: int, hidden: bool = False) -> SheetView:
    workbook = _workbook(path)
    if not 1 <= number <= len(workbook.worksheets):
        raise ValueError("This workbook has no such sheet.")
    ws = workbook.worksheets[number - 1]
    theme = _theme(workbook.loaded_theme)
    cells = ws._cells  # only the cells the file holds: ws.iter_rows would create every cell of the used range

    # the area with something in it: values, and merged areas that start among them
    filled = [key for key, cell in cells.items() if cell.value not in (None, "")]
    last_row = max((r for r, _ in filled), default=0)
    last_column = max((c for _, c in filled), default=0)
    merges = {}
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= last_row and merged.min_col <= last_column:
            merges[(merged.min_row, merged.min_col)] = merged
            last_row, last_column = max(last_row, merged.max_row), max(last_column, merged.max_col)

    widths, hidden_columns = _columns(ws, last_column)
    hidden_rows = {r for r in range(1, last_row + 1) if r in ws.row_dimensions and ws.row_dimensions[r].hidden}
    shown_columns = [c for c in range(1, last_column + 1) if hidden or c not in hidden_columns]
    shown_rows = [r for r in range(1, last_row + 1) if hidden or r not in hidden_rows]
    more_rows, more_columns = max(0, len(shown_rows) - MAX_ROWS), max(0, len(shown_columns) - MAX_COLUMNS)
    shown_rows, shown_columns = shown_rows[:MAX_ROWS], shown_columns[:MAX_COLUMNS]
    row_at = {r: i for i, r in enumerate(shown_rows)}
    column_at = {c: i for i, c in enumerate(shown_columns)}

    covered = {
        (r, c)
        for merged in merges.values()
        for r in range(merged.min_row, merged.max_row + 1)
        for c in range(merged.min_col, merged.max_col + 1)
    } - set(merges)
    styles: dict[Style, int] = {Style(): 0}
    rows = []
    default_height = ws.sheet_format.defaultRowHeight or DEFAULT_HEIGHT
    for r in shown_rows:
        out = []
        for c in shown_columns:
            if (r, c) in covered:
                continue
            cell = cells.get((r, c))
            merged = merges.get((r, c))
            style = _style(ws, cell, merged, theme) if cell is not None else Style()
            spans = (1, 1)
            if merged is not None:
                spans = (
                    sum(1 for x in range(r, merged.max_row + 1) if x in row_at),
                    sum(1 for x in range(c, merged.max_col + 1) if x in column_at),
                )
            out.append(
                Cell(
                    column=column_at[c],
                    text=display(cell.value, cell.number_format) if cell is not None else "",
                    style=styles.setdefault(style, len(styles)),
                    rows=max(spans[0], 1),
                    columns=max(spans[1], 1),
                )
            )
        height = ws.row_dimensions[r].ht if r in ws.row_dimensions else None
        rows.append(
            Row(number=r, height=round((height or default_height) * 4 / 3, 1), hidden=r in hidden_rows, cells=out)
        )

    frozen_rows = frozen_columns = 0
    if ws.freeze_panes:
        first_row, first_column = coordinate_to_tuple(ws.freeze_panes)  # the first cell that scrolls
        frozen_rows = sum(1 for r in shown_rows if r < first_row)
        frozen_columns = sum(1 for c in shown_columns if c < first_column)
    return SheetView(
        sheets=[SheetTab(name=s.title, hidden=s.sheet_state != "visible") for s in workbook.worksheets],
        number=number,
        right_to_left=bool(ws.sheet_view.rightToLeft),
        gridlines=ws.sheet_view.showGridLines is not False,
        frozen_rows=frozen_rows,
        frozen_columns=frozen_columns,
        columns=[
            Column(letter=get_column_letter(c), width=widths.get(c, _pixels(DEFAULT_WIDTH)), hidden=c in hidden_columns)
            for c in shown_columns
        ],
        rows=rows,
        styles=list(styles),
        more_rows=more_rows,
        more_columns=more_columns,
        hidden_rows=len(hidden_rows),
        hidden_columns=len(hidden_columns),
    )


def _pixels(characters: float) -> float:
    """A column width in characters of the default font, in pixels, as Excel works it out at 100%."""
    return float(int(characters * 7 + 5))


def _columns(ws, last: int) -> tuple[dict[int, float], set[int]]:
    """Each column's width in pixels, and the hidden columns. A width entry can cover a run of columns."""
    default = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.43
    widths = {c: _pixels(default) for c in range(1, last + 1)}
    hidden = set()
    for dimension in ws.column_dimensions.values():
        if dimension.min is None:
            continue
        for c in range(dimension.min, min(dimension.max or dimension.min, last) + 1):
            if dimension.width:
                widths[c] = _pixels(dimension.width)
            if dimension.hidden:
                hidden.add(c)
    return widths, hidden


def _style(ws, cell, merged, theme: list[str]) -> Style:
    font, fill, alignment, border = cell.font, cell.fill, cell.alignment, cell.border
    right, bottom = border.right, border.bottom
    if merged is not None:  # a merged area's outer edges belong to the cells along them
        right = getattr(ws._cells.get((merged.min_row, merged.max_col)), "border", border).right
        bottom = getattr(ws._cells.get((merged.max_row, merged.min_col)), "border", border).bottom
    align = ALIGN.get(alignment.horizontal or "")
    if align is None and isinstance(cell.value, (int, float, dt.date, dt.time)) and not isinstance(cell.value, bool):
        align = "end"  # Excel's general alignment puts numbers and dates at the end of the cell
    elif align is None and isinstance(cell.value, bool):
        align = "center"
    fill_colour = None
    if fill is not None and getattr(fill, "fill_type", None):  # a pattern shows as its main colour
        fill_colour = _colour(fill.fgColor, theme)
    elif fill is not None and getattr(fill, "stop", None):  # a gradient: its first colour
        fill_colour = _colour(fill.stop[0].color, theme)
    return Style(
        bold=bool(font.b),
        italic=bool(font.i),
        underline=bool(font.u),
        strike=bool(font.strike),
        font=font.name,
        size=float(font.sz) if font.sz else None,
        color=_colour(font.color, theme),
        fill=fill_colour,
        align=align,
        valign=VALIGN.get(alignment.vertical or ""),
        wrap=bool(alignment.wrap_text),
        indent=int(alignment.indent or 0),
        top=_side(border.top, theme),
        right=_side(right, theme),
        bottom=_side(bottom, theme),
        left=_side(border.left, theme),
    )


def _side(side, theme: list[str]) -> str | None:
    if side is None or not side.style:
        return None
    return f"{BORDERS.get(side.style, '1px solid')} {_colour(side.color, theme) or '#000000'}"


def _theme(xml: bytes | str | None) -> list[str]:
    """The workbook theme's colours, numbered as cells refer to them."""
    if not xml:
        return OFFICE_THEME
    try:
        root = ElementTree.fromstring(xml)
    except ElementTree.ParseError:
        return OFFICE_THEME
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    colours = []
    for name, fallback in zip(THEME, OFFICE_THEME, strict=True):
        node = scheme.find(f"a:{name}", ns) if scheme is not None else None
        value = None
        if node is not None and len(node):
            value = node[0].get("val") if node[0].tag.endswith("srgbClr") else node[0].get("lastClr")
        colours.append(value or fallback)
    return colours


def _colour(colour, theme: list[str]) -> str | None:
    if colour is None:
        return None
    if colour.type == "rgb" and isinstance(colour.rgb, str):
        value = colour.rgb[-6:]
    elif colour.type == "theme" and isinstance(colour.theme, int) and colour.theme < len(theme):
        value = theme[colour.theme]
    elif colour.type == "indexed" and isinstance(colour.indexed, int) and colour.indexed < 64:
        value = COLOR_INDEX[colour.indexed][-6:]  # 64 and 65 are the system's own colours: left to the screen
    else:
        return None
    return _tint(value, colour.tint) if colour.tint else f"#{value.lower()}"


def _tint(value: str, tint: float) -> str:
    """Excel's tint: towards black when negative, towards white when positive, by lightness."""
    red, green, blue = (int(value[i : i + 2], 16) / 255 for i in (0, 2, 4))
    hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
    lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
    return "#" + "".join(f"{round(v * 255):02x}" for v in colorsys.hls_to_rgb(hue, lightness, saturation))


# -- the text Excel shows -----------------------------------------------------------------------------------------


def display(value: object, number_format: str | None) -> str:
    """A cell's value as Excel displays it with its number format."""
    fmt = number_format or "General"
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return _date(value, fmt)
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        return _number(value, fmt)
    return str(value)


def _sections(fmt: str) -> list[str]:
    """A format's sections: positive; negative; zero; text. Semicolons inside quotes are text."""
    parts, current, quoted, escaped = [], "", False, False
    for ch in fmt:
        if escaped:
            current, escaped = current + ch, False
            continue
        if ch == "\\":
            escaped = True
        elif ch == '"':
            quoted = not quoted
        elif ch == ";" and not quoted:
            parts.append(current)
            current = ""
            continue
        current += ch
    return [*parts, current]


def _general(value: float) -> str:
    if isinstance(value, int) or value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return f"{value:.10g}".replace("e+", "E+").replace("e-", "E-")


def _number(value: float, fmt: str) -> str:
    sections = _sections(fmt)
    section, sign = sections[0], ""
    if value < 0 and len(sections) > 1:
        section, value = sections[1], -value
    elif value == 0 and len(sections) > 2:
        section = sections[2]
    elif value < 0:
        sign, value = "-", -value
    if section.strip().lower() in ("general", "", "@"):
        return sign + _general(value)
    # the section's parts: literal text, and the first run of digit placeholders, which the number replaces
    out, pattern, done, i = [], "", False, 0
    while i < len(section):
        ch = section[i]
        if ch == '"':
            end = section.find('"', i + 1)
            end = len(section) if end < 0 else end
            out.append(section[i + 1 : end])
            i = end + 1
            continue
        if ch == "[":  # a colour or condition, or [$SAR-401]: the symbol is shown
            end = section.find("]", i)
            end = len(section) if end < 0 else end
            inside = section[i + 1 : end]
            if inside.startswith("$"):
                out.append(inside[1:].split("-")[0])
            i = end + 1
            continue
        if ch == "\\" and i + 1 < len(section):
            out.append(section[i + 1])
            i += 2
            continue
        if ch == "_" and i + 1 < len(section):  # room for a character: a space
            out.append(" ")
            i += 2
            continue
        if ch == "*" and i + 1 < len(section):  # fill the cell with a character: nothing on screen
            i += 2
            continue
        if ch in "0#?.," and not done:
            start = i
            while i < len(section) and section[i] in "0#?.,":
                i += 1
            if not any(p in section[start:i] for p in "0#?"):
                out.append(section[start:i])
                continue
            pattern = section[start:i]
            if i < len(section) and section[i] in "eE" and section[i + 1 : i + 2] in ("+", "-"):
                places = pattern.partition(".")[2].count("0")
                out.append(f"{value:.{places}E}")
                i += 2
                while i < len(section) and section[i] in "0#?":
                    i += 1
            else:
                out.append("\0")  # where the number goes
            done = True
            continue
        if ch == "/" and done:  # a fraction: Excel's own way of showing it is beyond this
            return sign + _general(value)
        out.append(ch)
        i += 1
    if "%" in section:
        value *= 100 ** section.count("%")
    text = "".join(out)
    if "\0" in text:
        text = text.replace("\0", _digits(value, pattern))
    return sign + text.strip()


def _digits(value: float, pattern: str) -> str:
    integer, _, decimals = pattern.partition(".")
    scaled = len(integer) - len(integer.rstrip(","))  # each comma after the digits divides by a thousand
    integer = integer.rstrip(",")
    value /= 1000**scaled
    places = sum(ch in "0#?" for ch in decimals)
    required = len(decimals.rstrip("#?").replace(",", ""))
    whole, _, fraction = f"{value:.{places}f}".partition(".")
    while len(fraction) > required and fraction.endswith("0"):
        fraction = fraction[:-1]
    minimum = integer.count("0")
    whole = whole.zfill(minimum)
    if minimum == 0 and int(whole) == 0:
        whole = ""
    if "," in integer and whole:
        whole = f"{int(whole):,}".zfill(minimum)
    return whole + ("." + fraction if fraction or "." in pattern and places == 0 else "")


DATE_TOKENS = re.compile(
    r'"[^"]*"|\\.|\[[^\]]*\]|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|am/pm|a/p|\.0+|.', re.IGNORECASE
)


def _date(value: dt.date | dt.time, fmt: str) -> str:
    section = _sections(fmt)[0]
    if section.lower() in ("general", ""):
        return value.isoformat(sep=" ") if isinstance(value, dt.datetime) else value.isoformat()
    moment = value if isinstance(value, dt.datetime) else None
    if moment is None:
        moment = (
            dt.datetime.combine(value, dt.time())
            if isinstance(value, dt.date)
            else dt.datetime.combine(dt.date(1900, 1, 1), value)
        )
    tokens = DATE_TOKENS.findall(section)
    twelve = any(t.lower() in ("am/pm", "a/p") for t in tokens)
    out = []
    for n, token in enumerate(tokens):
        low = token.lower()
        if token.startswith('"'):
            out.append(token[1:-1])
        elif token.startswith("\\"):
            out.append(token[1:])
        elif token.startswith("["):
            if low in ("[h]", "[hh]"):
                out.append(str(moment.hour))
        elif low in ("m", "mm"):
            # minutes after an hour or before seconds; otherwise the month
            before = next(
                (t.lower().lstrip("[") for t in reversed(tokens[:n]) if t.lower().lstrip("[")[:1] in "hdys"), ""
            )
            after = next((t.lower() for t in tokens[n + 1 :] if t.lower()[:1] in "hdys"), "")
            minutes = before.startswith("h") or after.startswith("s")
            number = moment.minute if minutes else moment.month
            out.append(f"{number:02d}" if low == "mm" else str(number))
        elif low == "yyyy":
            out.append(f"{moment.year:04d}")
        elif low == "yy":
            out.append(f"{moment.year % 100:02d}")
        elif low == "mmmmm":
            out.append(moment.strftime("%B")[0])
        elif low == "mmmm":
            out.append(moment.strftime("%B"))
        elif low == "mmm":
            out.append(moment.strftime("%b"))
        elif low == "dddd":
            out.append(moment.strftime("%A"))
        elif low == "ddd":
            out.append(moment.strftime("%a"))
        elif low in ("dd", "d"):
            out.append(f"{moment.day:02d}" if low == "dd" else str(moment.day))
        elif low in ("hh", "h"):
            hour = (moment.hour % 12 or 12) if twelve else moment.hour
            out.append(f"{hour:02d}" if low == "hh" else str(hour))
        elif low in ("ss", "s"):
            out.append(f"{moment.second:02d}" if low == "ss" else str(moment.second))
        elif low == "am/pm":
            out.append("AM" if moment.hour < 12 else "PM")
        elif low == "a/p":
            out.append("A" if moment.hour < 12 else "P")
        elif low.startswith(".0"):
            out.append("." + f"{moment.microsecond:06d}"[: len(token) - 1])
        elif token == "_":
            out.append(" ")
        elif token not in ("*",):
            out.append(token)
    return "".join(out).strip()
