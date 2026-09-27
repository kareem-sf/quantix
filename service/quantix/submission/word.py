"""Documents as Word files the engineer can still edit: the firm's letterhead, a title block, real headings, lists and
tables, and page numbers. One quiet house style: Arial, dark ink, thin grey rules, no theme colours."""

from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph

from quantix.submission.content import (
    Document,
    Heading,
    Letterhead,
    ListBlock,
    Paragraph,
    Run,
    Table,
    is_arabic,
    plain,
)

FONT = "Arial"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
RULE = "C9CDD3"
SHADE = "F1F2F4"
SIZES = {"Normal": 10, "Title": 18, "Heading 1": 13, "Heading 2": 11.5, "Heading 3": 10.5}


def write(document: Document, head: Letterhead, path: Path, landscape: bool = False) -> None:
    word = docx.Document()
    _styles(word)
    section = word.sections[0]
    section.page_width, section.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin, section.bottom_margin = Cm(2.6), Cm(2)
    section.header_distance = section.footer_distance = Cm(1)
    width = (section.page_width - section.left_margin - section.right_margin) / 360000  # in cm
    _header(section, head, width)
    _footer(section, head, document.title, width)

    word.add_paragraph(document.title, style="Title")
    if document.facts:
        _facts(word, document.facts, width)
    for block in document.blocks:
        if isinstance(block, Heading):
            word.add_paragraph(block.text, style=f"Heading {block.level}")
        elif isinstance(block, Paragraph):
            _runs(word.add_paragraph(), block.runs)
        elif isinstance(block, ListBlock):
            _list(word, block)
        elif isinstance(block, Table):
            _table(word, block, width)
    word.save(path)


def _font(element, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    """Arial for every script, in place of the template's theme fonts, which Word would otherwise prefer."""
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in list(fonts.attrib):
        if attribute.endswith("Theme"):
            del fonts.attrib[attribute]
    for script in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(script), FONT)
    if size is not None:
        for tag in ("w:sz", "w:szCs"):
            node = rpr.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                rpr.append(node)
            node.set(qn("w:val"), str(int(size * 2)))
    if bold is not None:
        for tag in ("w:b", "w:bCs"):
            node = rpr.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                rpr.append(node)
            node.set(qn("w:val"), "1" if bold else "0")
    if color is not None:
        node = rpr.find(qn("w:color"))
        if node is None:
            node = OxmlElement("w:color")
            rpr.append(node)
        for attribute in list(node.attrib):
            del node.attrib[attribute]
        node.set(qn("w:val"), str(color))


def _styles(word) -> None:
    for name, size in SIZES.items():
        style = word.styles[name]
        _font(style.element, size, bold=name != "Normal", color=INK)
        paragraph = style.paragraph_format
        paragraph.space_before = Pt({"Normal": 0, "Title": 0, "Heading 1": 14, "Heading 2": 10}.get(name, 8))
        paragraph.space_after = Pt(4 if name.startswith("Heading") else 6)
        paragraph.line_spacing = 1.15
        borders = style.element.pPr.find(qn("w:pBdr")) if style.element.pPr is not None else None
        if borders is not None:  # the template underlines its Title
            style.element.pPr.remove(borders)
    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        word.styles[name].paragraph_format.space_after = Pt(3)


def _direction(paragraph: WordParagraph, text: str) -> None:
    if is_arabic(text):
        bidi = OxmlElement("w:bidi")
        paragraph._p.get_or_add_pPr().insert(0, bidi)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _runs(paragraph: WordParagraph, runs: list[Run], size: float | None = None, color: RGBColor | None = None):
    _direction(paragraph, plain(runs))
    for run in runs:
        pieces = run.text.split("\n")
        for index, piece in enumerate(pieces):
            if index:
                paragraph.add_run().add_break()
            if piece:
                word_run = paragraph.add_run(piece)
                word_run.bold = run.bold or None
                if size is not None:
                    word_run.font.size = Pt(size)
                if color is not None:
                    word_run.font.color.rgb = color
                if is_arabic(piece):
                    word_run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    return paragraph


def _borders(table: WordTable, sides: tuple[str, ...]) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single" if side in sides else "nil")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), RULE)
        borders.append(edge)
    properties.append(borders)


def _shade(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _cell(cell, text: str, bold: bool = False, right: bool = False, size: float = 9, color=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    _runs(paragraph, [Run(text, bold)], size=size, color=color)
    if right and not is_arabic(text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _widths(table: WordTable, shares: list[float], width: float) -> None:
    """Fixed column widths, which Word keeps only when the table, its grid and every cell all state them."""
    table.autofit = False
    total = sum(shares)
    for column, share in zip(table.columns, shares, strict=False):
        column.width = Cm(width * share / total)
    for row in table.rows:
        for cell, share in zip(row.cells, shares, strict=False):
            cell.width = Cm(width * share / total)
    margins = OxmlElement("w:tblCellMar")
    for side, size in (("top", 40), ("bottom", 40), ("left", 80), ("right", 80)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(size))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    table._tbl.tblPr.append(margins)


def _facts(word, facts: list[tuple[str, str]], width: float) -> None:
    table = word.add_table(rows=len(facts), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _borders(table, ("top", "bottom", "insideH"))
    for row, (label, value) in zip(table.rows, facts, strict=True):
        _cell(row.cells[0], label, color=MUTED)
        _cell(row.cells[1], value)
        _shade(row.cells[0], SHADE)
    _widths(table, [1, 3], width)
    word.add_paragraph()


def _table(word, block: Table, width: float) -> None:
    numeric = block.numeric
    table = word.add_table(rows=1 + len(block.rows) + len(block.totals), cols=len(block.header))
    _borders(table, ("top", "bottom", "insideH"))
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # repeats on every page
    for column, (cell, text) in enumerate(zip(header.cells, block.header, strict=True)):
        _cell(cell, text, bold=True, right=column in numeric)
        _shade(cell, SHADE)
    for index, values in enumerate(block.rows + block.totals, start=1):
        bold = index > len(block.rows)
        for column, (cell, text) in enumerate(zip(table.rows[index].cells, values, strict=False)):
            _cell(cell, text, bold=bold, right=column in numeric)
    _widths(table, block.widths or [1.0] * len(block.header), width)
    word.add_paragraph()


def _list(word, block: ListBlock) -> None:
    numbering = word.part.numbering_part.numbering_definitions._numbering
    restarted: dict[str, int] = {}  # one fresh count per numbered list, so each starts at 1
    for item in block.items:
        style = ("List Number" if item.numbered else "List Bullet") + (" 2" if item.level else "")
        paragraph = _runs(word.add_paragraph(style=style), item.runs)
        if item.numbered:
            if style not in restarted:
                style_num = word.styles[style].element.pPr.numPr.numId.val
                num = numbering.add_num(numbering.num_having_numId(style_num).abstractNumId.val)
                num.add_lvlOverride(ilvl=0).add_startOverride(1)
                restarted[style] = num.numId
            properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            properties.get_or_add_ilvl().val = 0
            properties.get_or_add_numId().val = restarted[style]


def _field(paragraph: WordParagraph, instruction: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    properties.append(size)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _header(section, head: Letterhead, width: float) -> None:
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(width))
    _borders(table, ("bottom",))
    left, right = table.rows[0].cells
    _cell(left, head.firm or head.tender, bold=True, size=9.5)
    for line in head.lines:
        paragraph = left.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        _runs(paragraph, [Run(line)], size=8, color=MUTED)
    if head.logo is not None:
        paragraph = right.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.add_run().add_picture(str(head.logo), height=Cm(1.3))
    _widths(table, [3, 1], width)
    header.paragraphs[0]._p.getparent().remove(header.paragraphs[0]._p)  # the template's empty first line


def _footer(section, head: Letterhead, title: str, width: float) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT)
    _runs(paragraph, [Run(f"{head.tender} · {title}\t")], size=8, color=MUTED)
    for text, instruction in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        run = paragraph.add_run(text)
        run.font.size, run.font.color.rgb = Pt(8), MUTED
        _field(paragraph, instruction)
