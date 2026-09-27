"""The package's Excel workbooks, laid out for people who will check and print them: a title block, a styled header
that repeats on every printed page, formulas the client can follow, number formats, and A4 fitted to the width."""

from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from quantix.boq.models import BoqItem
from quantix.estimate import records as estimate
from quantix.submission import package, records
from quantix.submission.content import Letterhead

FONT = "Arial"
INK, MUTED = "1F2937", "6B7280"
SHADE = PatternFill("solid", fgColor="F1F2F4")
BAND = PatternFill("solid", fgColor="E6E8EB")
RULE = Side(style="thin", color="C9CDD3")
LIGHT = Side(style="thin", color="E3E5E8")  # between lines; Excel prints "hair" nearly black
MONEY = "#,##0.00"


def _font(size: float = 10, bold: bool = False, color: str = INK) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color)


def _sheet(
    workbook: openpyxl.Workbook,
    title: str,
    head: Letterhead,
    facts: list[tuple[str, str]],
    widths: list[int],
    landscape: bool = False,
):
    """A sheet with the title block, the columns sized, and the letterhead and page numbers for printing."""
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.sheet_view.showGridLines = False
    sheet.cell(1, 1, title).font = _font(14, bold=True)
    for row, (label, value) in enumerate(facts, start=2):
        sheet.cell(row, 1, label).font = _font(9, color=MUTED)
        sheet.cell(row, 2, value).font = _font(9)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    # Excel reads "&" in headers and footers as a code; "&&" prints one
    sheet.oddHeader.left.text = (head.firm or head.tender).replace("&", "&&")
    sheet.oddHeader.left.size = 8
    sheet.oddFooter.left.text = f"{head.tender} · {title}".replace("&", "&&")
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.oddFooter.right.size = 8
    return sheet


def _header(sheet: Worksheet, row: int, labels: list[str], right: set[int]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row, column, label)
        cell.font, cell.fill = _font(9, bold=True), SHADE
        cell.border = Border(top=RULE, bottom=RULE)
        cell.alignment = Alignment(horizontal="right" if column in right else "left", vertical="center")
    sheet.freeze_panes = sheet.cell(row + 1, 1)
    sheet.print_title_rows = f"{row}:{row}"


def priced(
    session: Session,
    items: list[BoqItem],
    rates: dict[str, Decimal],
    summary: estimate.Summary,
    head: Letterhead,
    facts: list[tuple[str, str]],
    spread: bool,
    path: Path,
) -> None:
    """The priced BOQ in Quantix's layout, for bills with no rate column of their own: bill by bill in the client's
    row order, each amount the rounded product of its quantity and rate, subtotals, and the summary with VAT."""
    workbook = openpyxl.Workbook()
    sheet = _sheet(workbook, "Priced BOQ", head, facts, [7, 10, 58, 8, 13, 13, 16])
    row = len(facts) + 2
    sheet.cell(row, 1, package.basis(summary, spread)).font = _font(9, color=MUTED)
    row += 2
    _header(sheet, row, ["Row", "Item", "Description", "Unit", "Quantity", "Rate", "Amount"], {5, 6, 7})
    subtotals: list[tuple[str, int]] = []
    for name, lines in package.grouped(session, items):
        row += 1
        sheet.cell(row, 1, name).font = _font(10, bold=True)
        for column in range(1, 8):
            sheet.cell(row, column).fill = BAND
        first = row + 1
        for item in lines:
            row += 1
            rate = rates.get(item.id)
            values = [records.row_of(item.quote), item.item or None, item.description, item.unit, item.quantity, rate]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column, value)
                cell.font = _font(9)
                cell.alignment = Alignment(vertical="top", wrap_text=column == 3)
                cell.border = Border(bottom=LIGHT)
            sheet.cell(row, 5).number_format = "#,##0.00"
            sheet.cell(row, 6).number_format = MONEY
            amount = sheet.cell(row, 7, f"=ROUND(E{row}*F{row},2)" if rate is not None else None)
            amount.font, amount.number_format = _font(9), MONEY
            amount.alignment = Alignment(vertical="top")
            amount.border = Border(bottom=LIGHT)
        row += 1
        sheet.cell(row, 3, f"Subtotal, {name}").font = _font(9, bold=True)
        subtotal = sheet.cell(row, 7, f"=SUM(G{first}:G{row - 1})")
        subtotal.font, subtotal.number_format, subtotal.border = _font(9, bold=True), MONEY, Border(top=RULE)
        subtotals.append((name, row))

    row += 2
    sheet.cell(row, 1, "Summary").font = _font(11, bold=True)
    first = row + 1
    for name, at in subtotals:
        row += 1
        sheet.cell(row, 3, name).font = _font(9)
        cell = sheet.cell(row, 7, f"=G{at}")
        cell.font, cell.number_format = _font(9), MONEY
    if not spread and summary.total != summary.net:
        row += 1
        sheet.cell(row, 3, "Preliminaries, overheads and profit").font = _font(9)
        cell = sheet.cell(row, 7, summary.total - summary.net)
        cell.font, cell.number_format = _font(9), MONEY
    lines = [("Total before VAT" if summary.vat_rate is not None else "Total", f"=SUM(G{first}:G{row})")]
    total_row = row + 1
    if summary.vat_rate is not None:
        lines += [
            (package.vat_label(summary.vat_rate), f"=ROUND(G{total_row}*{summary.vat_rate},2)"),
            ("Total with VAT", f"=G{total_row}+G{total_row + 1}"),
        ]
    for label, formula in lines:
        row += 1
        sheet.cell(row, 3, label).font = _font(10, bold=True)
        cell = sheet.cell(row, 7, formula)
        cell.font, cell.number_format, cell.border = _font(10, bold=True), MONEY, Border(top=RULE)
    workbook.calculation.fullCalcOnLoad = True  # Excel works the formulas out as it opens the file
    workbook.save(path)


def checklist(rows: list[list[str | None]], head: Letterhead, facts: list[tuple[str, str]], path: Path) -> None:
    """What the tender requires and where each item stands, for the engineer; it is never sent."""
    workbook = openpyxl.Workbook()
    sheet = _sheet(workbook, "Checklist", head, facts, [16, 48, 34, 30, 40], landscape=True)
    row = len(facts) + 3
    _header(sheet, row, ["Section", "Requirement", "Required by", "State", "File"], set())
    for values in rows:
        row += 1
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.font = _font(9, color=INK if column != 4 or value == "Ready" else "C2410C")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=LIGHT)
    workbook.save(path)
