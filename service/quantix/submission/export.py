"""Build the submission package on the engineer's computer. Nothing is ever sent to the client."""

import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import docx
import openpyxl
from sqlalchemy.orm import Session

from quantix import tenders
from quantix.boq import records as boq
from quantix.boq.models import BoqItem
from quantix.documents import library
from quantix.documents.models import Document
from quantix.estimate import records as estimate
from quantix.submission import records


@dataclass
class Built:
    folder: str
    files: list[str] = field(default_factory=list)
    priced_total: Decimal = Decimal(0)
    summary_total: Decimal = Decimal(0)
    factor: Decimal = Decimal(1)
    not_ready: list[str] = field(default_factory=list)


def exports_dir(home: Path) -> Path:
    return home / "exports"


def _safe(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", name).strip()[:120] or "Tender"


def submitted_rates(session: Session, tender_id: str, spread: bool) -> tuple[dict[str, Decimal], Decimal]:
    """The rate to write for each priced item. Spreading lifts every rate by total ÷ net, so the markups sit in the
    rates and the BOQ adds up to the tender total, give or take the rounding of each rate to the cent."""
    summary = estimate.summary(session, tender_id)
    factor = summary.total / summary.net if spread and summary.net else Decimal(1)
    rates = {}
    for item in boq.items(session, tender_id):
        rate = estimate.current_rate(session, item.id)
        if rate is not None and item.quantity is not None:
            rates[item.id] = estimate.money(estimate.rate_of(rate) * factor)
    return rates, factor


def _client_format(home: Path, session: Session, tender_id: str, folder: Path, rates: dict[str, Decimal]):
    """Copies of the client's BOQ workbooks with rates and amounts written into their own columns. Returns the files
    and the items those sheets cover, priced or not."""
    covered: set[str] = set()
    files = []
    items = boq.items(session, tender_id)
    columns = records.pricing_columns(session, tender_id)
    for document_id in {c.document_id for c in columns}:
        document = session.get(Document, document_id)
        source = library.stored_file(home, document)
        workbook = openpyxl.load_workbook(source, keep_vba=source.suffix == ".xlsm")
        for sheet in (c for c in columns if c.document_id == document_id):
            worksheet = workbook.worksheets[sheet.sheet - 1]
            for item in items:
                row = records.row_of(item.quote)
                if item.document_id != document_id or item.page != sheet.sheet or row is None:
                    continue
                covered.add(item.id)
                if item.id not in rates:
                    continue
                worksheet[f"{sheet.rate_column}{row}"] = rates[item.id]
                amount = worksheet[f"{sheet.amount_column}{row}"]
                if not (isinstance(amount.value, str) and amount.value.startswith("=")):  # keep the client's formula
                    amount.value = estimate.money(item.quantity * rates[item.id])
        name = f"Priced {document.name}"
        workbook.save(folder / name)
        files.append(name)
    return files, covered


def _quantix_format(folder: Path, items: list[BoqItem], rates: dict[str, Decimal], currency: str) -> str:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Priced BOQ"
    sheet.append(["Item", "Description", "Unit", "Quantity", f"Rate {currency}".strip(), f"Amount {currency}".strip()])
    for item in items:
        rate = rates.get(item.id)
        amount = estimate.money(item.quantity * rate) if rate is not None else None
        sheet.append([item.item, item.description, item.unit, item.quantity, rate, amount])
    workbook.save(folder / "Priced BOQ.xlsx")
    return "Priced BOQ.xlsx"


def _draft_file(folder: Path, title: str, body: str) -> str:
    document = docx.Document()
    document.add_heading(title, level=1)
    for paragraph in re.split(r"\n\s*\n", body):
        document.add_paragraph(paragraph.strip())
    name = f"{_safe(title)}.docx"
    document.save(folder / name)
    return name


def build(home: Path, session: Session, tender_id: str, spread: bool, now: datetime) -> Built:
    tender = tenders.get_tender(session, tender_id)
    folder = exports_dir(home) / f"{_safe(tender.name)} {now:%Y-%m-%d %H%M}"
    folder.mkdir(parents=True, exist_ok=True)
    built = Built(folder=folder.name)

    items = boq.items(session, tender_id)
    rates, built.factor = submitted_rates(session, tender_id, spread)
    summary = estimate.summary(session, tender_id)
    built.summary_total = summary.total
    built.priced_total = sum((estimate.money(i.quantity * rates[i.id]) for i in items if i.id in rates), Decimal(0))
    client_files, covered = _client_format(home, session, tender_id, folder, rates)
    built.files += client_files
    if any(i.id not in covered for i in items):
        built.files.append(_quantix_format(folder, [i for i in items if i.id not in covered], rates, summary.currency))
    built.not_ready += [f"BOQ item {i} is not priced" for i in summary.unpriced]
    if summary.waiting:
        built.not_ready.append(f"{summary.waiting} rates still wait for your approval")

    checklist = openpyxl.Workbook()
    sheet = checklist.active
    sheet.title = "Checklist"
    sheet.append(["Section", "Requirement", "Required by", "State", "File"])
    for requirement in records.requirements(session, tender_id):
        state, file = records.state(session, requirement), None
        current = records.current_draft(session, requirement.id)
        if requirement.file_name:
            file = f"{_safe(requirement.title)} - {requirement.file_name}"
            shutil.copyfile(records.attachments_dir(home, requirement) / requirement.file_name, folder / file)
        elif current is not None and current.status != "proposed":
            file = _draft_file(folder, current.title, current.body)
        if file:
            built.files.append(file)
        if state != "ready":
            built.not_ready.append(requirement.title)
        source = session.get(Document, requirement.document_id) if requirement.document_id else None
        label = {"ready": "Ready", "review": "Draft waiting for review", "missing": "Missing"}[state]
        if current is not None and current.status == "office_approved" and not requirement.file_name:
            label = "Ready · approved by the office, not reviewed"
        sheet.append(
            [
                requirement.section,
                requirement.title,
                f"{source.name}, page {requirement.page}" if source else "Added by the engineer",
                label,
                file,
            ]
        )
    checklist.save(folder / "Checklist.xlsx")
    built.files.append("Checklist.xlsx")
    return built
