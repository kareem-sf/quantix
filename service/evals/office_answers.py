"""Asks the office's own AI how a synthetic tender's lines were priced, and checks each answer: it came in the same
turn, it rests on the rate's record, and it took few steps. It uses the office AI chosen in Settings (read from
~/.quantix, never copied) on a synthetic tender in a scratch data home, and costs a few cents.

Run from service/:  .venv/Scripts/python evals/office_answers.py
"""

import io
import shutil
import sys
import tempfile
import time
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import docx
import openpyxl
from fastapi.testclient import TestClient
from pydantic_evals import Case, Dataset
from pydantic_evals.evaluators import Evaluator, EvaluatorContext

from quantix import settings
from quantix.api.app import create_app
from quantix.boq import records as boq
from quantix.estimate import records as estimate
from quantix.office import records as office
from quantix.office import runtime
from quantix.office.models import Task, TurnRecord
from quantix.review import records as reviews

TOKEN = "eval"
ALLOWANCE = 600_000  # tokens the synthetic tender may use, so a confused office can't run on
NOTE = "Fixed by a gang of four at 16 hours a tonne; rebar at the supplier price with 5% wastage and tie wire."
BUILD_UP = [
    estimate.LineIn(kind="labour", resource="Steel fixer gang", quantity=Decimal(16), unit="hr", rate=Decimal(62)),
    estimate.LineIn(kind="material", resource="Rebar B500B cut and bent", quantity=Decimal(1), unit="t",
                    rate=Decimal(2300), wastage=Decimal("0.05")),
    estimate.LineIn(kind="material", resource="Tie wire", quantity=Decimal(12), unit="kg", rate=Decimal(6)),
]  # fmt: skip


def _bill() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BOQ"
    for row in (
        ["Item", "Description", "Unit", "Qty"],
        ["3.1", "Excavation to reduce levels", "m3", 1240],
        ["4.3", "Slab reinforcement", "t", 28.1],
    ):
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _conditions() -> bytes:
    document = docx.Document()
    document.add_paragraph("Prices shall be in Saudi Riyals (SAR). VAT at 15% shall be shown separately.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _synthetic(client: TestClient) -> tuple[str, str]:
    """A school tender whose slab reinforcement was sent back once, then priced from a build-up and approved."""
    tender_id = client.post("/tenders", json={"name": "Synthetic school"}).json()["id"]
    files = [("files", ("Bill.xlsx", _bill(), "")), ("files", ("Conditions.docx", _conditions(), ""))]
    client.post(f"/tenders/{tender_id}/documents", files=files)
    for _ in range(200):
        documents = client.get(f"/tenders/{tender_id}/documents").json()
        if all(d["status"] == "read" for d in documents):
            break
        time.sleep(0.1)
    bill = next(d["id"] for d in documents if d["name"] == "Bill.xlsx")
    with client.app.state.sessions() as session:
        rania = office.hire(session, tender_id, "Rania Farouk", "Tender Manager", {}, is_manager=True)
        priya = office.hire(session, tender_id, "Priya Nair", "Estimator", {"work": ["boq", "pricing"]})
        rows = [("3.1", "Excavation to reduce levels", "m3", "1240"), ("4.3", "Slab reinforcement", "t", "28.1")]
        lines = [
            boq.ItemIn(item=i, description=d, unit=u, quantity=Decimal(q), document_id=bill, page=1,
                       quote=f"A{n}={i} | B{n}={d} | C{n}={u} | D{n}={q}")
            for n, (i, d, u, q) in enumerate(rows, start=2)
        ]  # fmt: skip
        boq.propose_items(session, tender_id, priya, lines)
        for item in boq.items(session, tender_id):
            boq.approve(session, item)
        excavation = estimate.propose_rate(session, tender_id, priya.id, "3.1", "estimate",
                                           "Excavator at 45 m3 an hour, 830 SAR an hour with driver.",
                                           unit_rate=Decimal("18.50"))  # fmt: skip
        estimate.approve(session, excavation)
        estimate.propose_rate(session, tender_id, priya.id, "4.3", "estimate", "Rebar at the supplier price only.",
                              unit_rate=Decimal(2300))  # fmt: skip
        [first] = reviews.pending(session, tender_id)
        sent_back = reviews.Verdict(record=first.ref, accept=False, note="Add the fixing labour and the tie wire.")
        reviews.review(session, client.app.state.home, tender_id, rania, [sent_back], autonomous=False)
        rebar = estimate.propose_rate(session, tender_id, priya.id, "4.3", "estimate", NOTE, lines=BUILD_UP)
        [second] = reviews.pending(session, tender_id)
        accepted = reviews.Verdict(record=second.ref, accept=True, note="Checked the gang output and the wastage.")
        reviews.review(session, client.app.state.home, tender_id, rania, [accepted], autonomous=False)
        estimate.decide(session, rebar, True)
        for task in session.query(Task).filter(Task.status == "open"):
            task.status = "done"  # Priya's redo is done: nothing but the engineer's question is new
        session.commit()
        return tender_id, rania.id


@dataclass
class Answer:
    text: str = ""
    sources: list[str] = field(default_factory=list)
    follow_ups: list[str] = field(default_factory=list)
    requests: int = 0  # model requests in the Manager's turn that answered
    tokens: int = 0  # in every turn the question led to
    turns: int = 0
    calls: list[str] = field(default_factory=list)  # in the turn that answered


def ask(question: str) -> Answer:
    """One question to the Tender Manager on a fresh synthetic tender, and what came back."""
    model = runtime.office_model(Path.home() / ".quantix")
    if model is None:
        sys.exit("Choose the office's AI in Settings first.")
    home = Path(tempfile.mkdtemp(prefix="quantix-eval-"))
    try:
        with TestClient(create_app(home, TOKEN), headers={"Authorization": f"Bearer {TOKEN}"}) as client:
            settings.save(home, office_ai={"connection_id": "eval", "model": model.model_name},
                          tender_allowance=ALLOWANCE)  # fmt: skip
            client.app.state.office.model = lambda: model
            client.app.state.office.sees_images = lambda: False
            tender_id, rania = _synthetic(client)
            client.post(f"/tenders/{tender_id}/messages", json={"channel": rania, "text": question})
            deadline, quiet = time.monotonic() + 300, 0
            while time.monotonic() < deadline and quiet < 3:
                time.sleep(2)
                quiet = quiet + 1 if client.get(f"/tenders/{tender_id}/office").json()["state"] != "working" else 0
            client.post(f"/tenders/{tender_id}/office/stop")
            chat = client.get(f"/tenders/{tender_id}/messages", params={"channel": rania}).json()
            answer = Answer()
            reply = next((m for m in chat if m["sender"] == rania), None)  # the reply to the question
            if reply is not None:
                answer.text = reply["text"]
                answer.sources = [s["label"] for s in reply["sources"] or []]
            with client.app.state.sessions() as session:
                answer.follow_ups = [t.title for t in session.query(Task).filter(Task.staff_id == rania)]
                turns = session.query(TurnRecord).order_by(TurnRecord.id).all()
                first = next((t for t in turns if t.staff_id == rania), None)
                answer.requests = first.requests if first else 0  # in the turn that answered
                answer.tokens = sum(t.input_tokens + t.output_tokens for t in turns)
                answer.turns = len(turns)
                answer.calls = [c["tool"] for c in first.calls] if first else []
            return answer
    finally:
        shutil.rmtree(home, ignore_errors=True)


@dataclass
class Answered(Evaluator[str, Answer]):
    """The Manager wrote back."""

    def evaluate(self, ctx: EvaluatorContext[str, Answer]) -> bool:
        return bool(ctx.output.text)


@dataclass
class RestsOnTheRate(Evaluator[str, Answer]):
    """The answer names the line's rate or BOQ line among its sources."""

    line: str = "4.3"

    def evaluate(self, ctx: EvaluatorContext[str, Answer]) -> bool:
        return any(self.line in source for source in ctx.output.sources)


@dataclass
class NoBarePromise(Evaluator[str, Answer]):
    """An answer without sources had its remaining work set as follow-up tasks."""

    def evaluate(self, ctx: EvaluatorContext[str, Answer]) -> bool:
        return bool(ctx.output.sources or ctx.output.follow_ups)


@dataclass
class FewSteps(Evaluator[str, Answer]):
    most: int = 12  # one turn's allowance: the answer came in the turn that read the question

    def evaluate(self, ctx: EvaluatorContext[str, Answer]) -> bool:
        return ctx.output.requests <= self.most


QUESTIONS = Dataset[str, Answer](
    name="the office answers from its records",
    cases=[
        Case(name="how each item was priced", inputs="How did you price each item?"),
        Case(name="the source of a rate", inputs="What is the source for the slab reinforcement rate?"),
        Case(name="one line in detail", inputs="Regarding the slab reinforcement, give me the details for pricing it."),
    ],
    evaluators=[Answered(), RestsOnTheRate(), NoBarePromise(), FewSteps()],
)


if __name__ == "__main__":
    report = QUESTIONS.evaluate_sync(ask, max_concurrency=1)
    report.print(include_input=True, include_output=True, include_durations=False)
