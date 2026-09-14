"""Auditable consolidated BOQ and Tender-record workbooks."""

import math
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _excel_text(value):
    """Imported descriptions and formulas are inert text in generated workbooks."""
    return str(value or "")[:32700]


def _write_text(cell, value):
    cell.value = _excel_text(value)
    cell.data_type = "s"


def _number(value):
    return float(Decimal(value)) if value is not None else None


def _sheet(book, name, tender, headers, widths):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C7"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _write_text(sheet.cell(1, 1), tender)
    sheet.cell(1, 1).font = Font(name="Aptos", size=18, bold=True, color="173E38")
    sheet.row_dimensions[1].height = 32
    _write_text(sheet.cell(2, 1), "Draft pending engineer review and final release")
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        _write_text(sheet.cell(6, column), header)
        sheet.cell(6, column).fill = PatternFill("solid", fgColor="173E38")
        sheet.cell(6, column).font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        sheet.cell(6, column).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[6].height = 32
    sheet.print_title_rows = "1:6"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def _cell_chunks(value, width):
    chunks, start, column, lines = [], 0, 0, 1
    for index, character in enumerate(value):
        column += 1
        if character == "\n" or column >= max(1, width - 3):
            lines += 1
            column = 0
        if lines >= 20 or index - start >= 799:
            chunks.append(value[start : index + 1])
            start, column, lines = index + 1, 0, 1
    if start < len(value):
        chunks.append(value[start:])
    return chunks or [""]


def _tab(book, name, tender, headers, widths, rows):
    sheet = _sheet(book, name, tender, headers, widths)
    _write_text(sheet.cell(3, 1), name)
    sheet.cell(3, 1).font = Font(name="Aptos", size=13, bold=True, color="173E38")
    sheet.row_dimensions[6].height = 48
    row_number = 7
    for values in rows:
        # Long saved replies and source passages continue on additional rows, never silently truncate.
        parts = [
            _cell_chunks(value, widths[index]) if isinstance(value, str) else [value]
            for index, value in enumerate(values)
        ]
        for part in range(max(map(len, parts))):
            for column, chunks in enumerate(parts, 1):
                value = chunks[part] if part < len(chunks) else ""
                cell = sheet.cell(row_number, column)
                if isinstance(value, (int, float)):
                    cell.value = value
                else:
                    _write_text(cell, value)
                cell.font = Font(name="Aptos", size=10, color="173E38")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if row_number % 2:
                    cell.fill = PatternFill("solid", fgColor="F0F5F3")
            lines = max(
                sum(
                    max(1, math.ceil(len(line) / max(1, width - 2)))
                    for line in str(sheet.cell(row_number, col).value or "").split("\n")
                )
                for col, width in enumerate(widths, 1)
            )
            sheet.row_dimensions[row_number].height = min(409, max(36, lines * 13 + 8))
            row_number += 1
    sheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}{max(6, row_number - 1)}"
    return sheet


def _extra_excel(path, tender, request, data, sources, warnings):
    book = Workbook()
    book.remove(book.active)
    _tab(
        book,
        "Review notes",
        tender["name"],
        ["Review point", "Detail"],
        [32, 110],
        [
            (
                "Scope",
                "Selected saved records only. Registered, extracted, analysed and reviewed coverage are distinct.",
            ),
            *(("Outstanding", warning) for warning in warnings),
        ],
    )
    if request.kind == "registers_xlsx":
        for name, kind in [
            ("Clarifications", "question"),
            ("Assumptions", "assumption"),
            ("Exclusions", "exclusion"),
            ("Risks", "risk"),
        ]:
            rows = [
                [
                    item["id"],
                    item["title"],
                    item["detail"]
                    + (
                        "\nRejected exclusion. This exclusion does not apply."
                        if kind == "exclusion" and item["state"] == "rejected"
                        else ""
                    ),
                    item["state"],
                    "Source changed" if item.get("is_stale") else "Current source record",
                    ", ".join(item["source_ids"]),
                ]
                for item in data["findings"]
                if item["kind"] == kind
            ]
            _tab(
                book,
                name,
                tender["name"],
                [
                    "Record ID",
                    "Subject",
                    "Detail",
                    "Engineer decision",
                    "Source state",
                    "Evidence IDs",
                ],
                [38, 40, 75, 22, 24, 38],
                rows or [["", f"No {name.lower()} recorded", "", "", "", ""]],
            )
    elif request.kind == "comparison_xlsx":
        _tab(
            book,
            "Supplier requests",
            tender["name"],
            [
                "Request ID",
                "Recipients",
                "Subject",
                "Request text",
                "Delivery state",
                "Evidence IDs",
            ],
            [25, 35, 40, 65, 24, 30],
            [
                [
                    q["id"],
                    ", ".join(q["to"]),
                    q["subject"],
                    q["body"],
                    q["status"],
                    ", ".join(q["source_ids"]),
                ]
                for q in data["quotes"]
            ]
            or [["", "No supplier requests recorded", "", "", "", ""]],
        )
        _tab(
            book,
            "Supplier replies",
            tender["name"],
            [
                "Reply ID",
                "Request ID",
                "Supplier",
                "Received",
                "Saved reply",
                "Evidence IDs",
                "Commercial decision",
            ],
            [24, 24, 32, 24, 65, 28, 25],
            [
                [
                    r["id"],
                    r["quote_id"],
                    r["sender"],
                    r["received_at"],
                    r["text"],
                    ", ".join(r["source_ids"] + r.get("supporting_source_ids", [])),
                    "Receipt is not acceptance",
                ]
                for r in data["replies"]
            ]
            or [["", "", "No supplier replies recorded", "", "", "", ""]],
        )
        items = {item["id"]: item for item in data["view"]["items"]}
        rows = []
        for proposed in data["rates"]:
            payload = proposed["payload"]
            provenance = payload["provenance"]
            matched = [
                r for r in data["replies"] if set(proposed["source_ids"]) & set(r["source_ids"])
            ]
            amount = payload["unit_rate"]
            if amount is None:
                amount = str(
                    sum(
                        (
                            Decimal(c["quantity"]) * Decimal(c["unit_rate"])
                            for c in payload["components"]
                        ),
                        Decimal(0),
                    )
                )
            item = items.get(proposed["item_id"], {})
            rows.append(
                [
                    proposed["id"],
                    item.get("description", proposed["item_id"]),
                    item.get("unit", "Unknown"),
                    ", ".join(sorted({r["sender"] for r in matched})) or "Unlinked rate proposal",
                    amount,
                    payload["currency"],
                    {
                        "including_vat": "Including VAT",
                        "excluding_vat": "Excluding VAT",
                        "unknown": "Unknown",
                    }[payload["tax_basis"]],
                    payload["vat_percent"] or "Unknown",
                    proposed["status"],
                    "Current" if proposed["is_current"] else "Source or estimate changed",
                    provenance["basis"],
                    provenance["observed_on"],
                    provenance["geography"],
                    provenance["conditions"],
                    ", ".join(proposed["source_ids"] + provenance["urls"]),
                ]
            )
        _tab(
            book,
            "Rate comparison",
            tender["name"],
            [
                "Proposal ID",
                "BOQ item",
                "Unit",
                "Linked supplier reply",
                "Rate as supplied",
                "Currency",
                "Tax basis",
                "VAT percent",
                "Engineer decision",
            ],
            [24, 42, 12, 32, 18, 14, 22, 16, 22],
            [row[:9] for row in rows] or [["No rate proposals recorded"] + [""] * 8],
        )
        _tab(
            book,
            "Rate provenance",
            tender["name"],
            [
                "Proposal ID",
                "Basis state",
                "Price basis",
                "Price date",
                "Geography",
                "Conditions",
                "Sources",
            ],
            [24, 28, 18, 18, 24, 65, 55],
            [[row[0], *row[9:]] for row in rows] or [["No rate proposals recorded"] + [""] * 6],
        )
    else:
        programme = request.programme
        _tab(
            book,
            "Programme",
            tender["name"],
            [
                "Activity ID",
                "Construction activity",
                "Working days",
                "Predecessors",
                "Start",
                "Finish",
                "Assumptions",
                "Evidence IDs",
            ],
            [20, 55, 18, 24, 18, 18, 60, 38],
            [
                [
                    r["id"],
                    r["title"],
                    r["duration_days"],
                    ", ".join(r["predecessor_ids"]),
                    r["start_date"],
                    r["finish_date"],
                    "\n".join(r["assumptions"]),
                    ", ".join(r["source_ids"]),
                ]
                for r in data["scheduled_activities"]
            ],
        )
        weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
        _tab(
            book,
            "Calendar",
            tender["name"],
            ["Calendar input", "Value"],
            [32, 110],
            [
                ("Programme", programme.title),
                ("Requested start", programme.start_date.isoformat()),
                ("Working week", ", ".join(weekdays[i] for i in programme.working_week)),
                (
                    "Holidays",
                    ", ".join(d.isoformat() for d in programme.holidays) or "None supplied",
                ),
                (
                    "Sequencing",
                    "Whole working days inclusive. Successors start on the next working day after all predecessors finish. No resource levelling or contractual completion assessment.",
                ),
                *(("Programme assumption", a) for a in programme.assumptions),
            ],
        )
    _tab(
        book,
        "Sources",
        tender["name"],
        [
            "Evidence ID",
            "Source document",
            "Exact location",
            "Version",
            "Source hash",
            "Source text",
        ],
        [28, 36, 32, 12, 40, 80],
        [
            [s["id"], s["relative_path"], s["locator"], s["version"], s["content_hash"], s["text"]]
            for s in sources
        ],
    )
    book.save(path)
    book.close()


def _excel(path, tender, view, sources):
    book = Workbook()
    book.remove(book.active)
    summary = _sheet(
        book,
        "Summary",
        tender["name"],
        [
            "Currency",
            "Priced subtotal excluding VAT",
            "Complete total excluding VAT",
            "Complete total including VAT",
            "Pricing state",
        ],
        [15, 25, 25, 25, 25],
    )
    _write_text(summary["A3"], "Pricing")
    _write_text(
        summary["B3"], "Complete for confirmed candidate rows" if view["complete"] else "Incomplete"
    )
    _write_text(summary["A4"], "Outstanding")
    _write_text(
        summary["B4"],
        " ".join(view["blocking_reasons"]) or "Final release remains an engineer decision.",
    )
    summary.merge_cells("B4:E4")
    summary["B4"].alignment = Alignment(wrap_text=True)
    summary.row_dimensions[4].height = 45
    for row, total in enumerate(view["totals"], 7):
        summary.append(
            [
                total["currency"],
                _number(total["priced_subtotal_ex_vat"]),
                _number(total["total_ex_vat"]),
                _number(total["total_inc_vat"]),
                "Complete" if total["complete"] else "Incomplete",
            ]
        )
    boq = _sheet(
        book,
        "BOQ",
        tender["name"],
        [
            "Description",
            "Unit",
            "Currency",
            "Supplied quantity",
            "Confirmed quantity used",
            "Rate excluding VAT",
            "Amount excluding VAT",
            "VAT",
            "Amount including VAT",
            "Quantity basis",
            "Source state",
            "Evidence ID",
            "Supplied unit rate",
            "Supplied rate tax basis",
        ],
        [52, 10, 12, 17, 19, 19, 20, 12, 20, 24, 22, 38, 19, 24],
    )
    _write_text(boq["A3"], view["coverage_note"])
    for row, item in enumerate(view["items"], 7):
        values = [
            item["description"],
            item["unit"],
            item["currency"],
            _number(item["supplied_quantity"]),
            _number(item["effective_quantity"]) if item["confirmed"] else None,
            _number(item["rate_ex_vat"]),
            f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(F{row})),ROUND(E{row}*F{row},2),"")',
            _number(item["vat_percent"]) / 100 if item["vat_percent"] is not None else None,
            f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(M{row}),ISNUMBER(H{row})),ROUND(E{row}*M{row},2),"")'
            if item["tax_basis"] == "including_vat"
            else f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(F{row}),ISNUMBER(H{row})),ROUND(E{row}*F{row}*(1+H{row}),2),"")',
            "Approved measurement"
            if item["quantity_basis"] == "approved_measurement"
            else "Supplied BOQ",
            "Confirmed" if item["confirmed"] else "Needs confirmation",
            item["source_id"],
            _number(item["unit_rate"]),
            {
                "including_vat": "Including VAT",
                "excluding_vat": "Excluding VAT",
                "unknown": "Unknown",
            }[item["tax_basis"]],
        ]
        for column, value in enumerate(values, 1):
            target = boq.cell(row, column)
            if isinstance(value, str) and column not in {7, 9}:
                _write_text(target, value)
            else:
                target.value = value
        boq.cell(row, 8).number_format = "0.00%"
    rates = _sheet(
        book,
        "Rate build-ups",
        tender["name"],
        [
            "BOQ description",
            "Component",
            "Consumption",
            "Component unit rate",
            "Rate contribution",
            "Unit",
            "Currency",
            "Tax basis",
            "Rate basis",
            "Source date",
            "Geography",
            "Conditions",
            "Source references",
        ],
        [42, 24, 16, 20, 20, 12, 12, 20, 16, 16, 22, 48, 50],
    )
    row = 7
    for item in view["items"]:
        components = item["components"] or (
            [
                {
                    "name": "Direct unit rate",
                    "quantity": "1",
                    "unit_rate": item["unit_rate"],
                    "unit": item["unit"],
                }
            ]
            if item["unit_rate"] is not None
            else []
        )
        source = item["provenance"] or {}
        for component in components:
            values = [
                item["description"],
                component["name"],
                _number(component["quantity"]),
                _number(component["unit_rate"]),
                f"=C{row}*D{row}",
                component["unit"],
                item["currency"],
                {
                    "excluding_vat": "Excluding VAT",
                    "including_vat": "Including VAT",
                    "unknown": "Unknown",
                }[item["tax_basis"]],
                source.get("basis"),
                source.get("observed_on"),
                source.get("geography"),
                source.get("conditions"),
                "\n".join(source.get("source_ids", []) + source.get("urls", [])),
            ]
            for column, value in enumerate(values, 1):
                if isinstance(value, str) and column != 5:
                    _write_text(rates.cell(row, column), value)
                else:
                    rates.cell(row, column).value = value
            row += 1
    source_sheet = _sheet(
        book,
        "Sources",
        tender["name"],
        ["Evidence ID", "Source document", "Exact location", "Extracted source text"],
        [38, 45, 38, 95],
    )
    for row, source in enumerate(sources, 7):
        for col, value in enumerate(
            [source["id"], source["relative_path"], source["locator"], source["text"]], 1
        ):
            _write_text(source_sheet.cell(row, col), value)
    for sheet in book:
        sheet.auto_filter.ref = f"A6:{get_column_letter(sheet.max_column)}{max(sheet.max_row, 6)}"
        for row in sheet.iter_rows(min_row=7):
            for cell in row:
                cell.font = Font(name="Aptos", size=10, color="173E38")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row % 2:
                    cell.fill = PatternFill("solid", fgColor="F0F5F3")
                if cell.data_type in {"n", "f"} and cell.number_format == "General":
                    cell.number_format = "#,##0.00;[Red](#,##0.00);–"
            line_count = max(
                math.ceil(
                    len(str(cell.value or ""))
                    / max(1, sheet.column_dimensions[cell.column_letter].width - 2)
                )
                for cell in row
            )
            sheet.row_dimensions[row[0].row].height = min(409, max(48, line_count * 13 + 8))
        sheet.print_options.horizontalCentered = True
    book.save(path)
    book.close()
