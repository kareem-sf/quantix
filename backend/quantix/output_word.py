"""Word engineering documents from saved Tender records."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _technical_word(path, tender, task, sources, warnings):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.add_paragraph("Technical work document", "Title")
    document.add_paragraph(tender["name"] + "\n" + task["title"])
    document.add_paragraph("Draft pending engineer review and final release")
    document.add_paragraph(
        f"Saved specialist work by {task['role']}. Task {task['id']}. This document covers the selected completed task only."
    )
    document.add_heading("Work scope", 1)
    document.add_paragraph(task["description"])
    document.add_heading("Saved specialist result", 1)
    result = task["result"]
    for paragraph in result["summary"].splitlines():
        if paragraph.strip():
            document.add_paragraph(paragraph)
    document.add_paragraph(
        "Result evidence: "
        + (", ".join(result.get("source_ids", [])) or "No source reference recorded")
    )
    for finding in result.get("findings", []):
        document.add_heading(finding["title"], 2)
        document.add_paragraph(f"{finding['kind']} — {finding.get('state', 'proposed')}")
        document.add_paragraph(finding["detail"])
        document.add_paragraph(
            "Evidence: "
            + (", ".join(finding.get("source_ids", [])) or "No source reference recorded")
        )
    for finding in result.get("web_findings", []):
        document.add_heading(finding["title"], 2)
        document.add_paragraph(finding["detail"])
        document.add_paragraph("Web sources: " + ", ".join(finding.get("urls", [])))
    document.add_heading("Missing information and review", 1)
    for warning in warnings:
        document.add_paragraph(warning)
    document.add_heading("Source references", 1)
    for source in sources:
        document.add_paragraph(
            f"{source['relative_path']} — {source['locator']}\nEvidence ID: {source['id']}\nVersion {source['version']} | SHA256 {source['content_hash']}"
        )
    section.footer.paragraphs[0].text = "Draft | "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    section.footer.paragraphs[0]._p.append(field)
    document.core_properties.author = "Quantix"
    document.core_properties.title = task["title"]
    document.save(path)


def _word(path, tender, overview, view, findings, messages, sources):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.add_paragraph("Tender analysis", "Title")
    document.add_paragraph(tender["name"])
    document.add_paragraph("Draft pending engineer review and final release")
    document.add_paragraph(
        "Pricing is "
        + (
            "complete for the confirmed candidate rows."
            if view["complete"]
            else "incomplete. " + " ".join(view["blocking_reasons"])
        )
    )
    document.add_heading("Tender Manager analysis", 1)
    latest = next((message for message in reversed(messages) if message["role"] == "manager"), None)
    if latest:
        for paragraph in latest["content"].split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph)
    else:
        document.add_paragraph("The Tender Manager has not yet recorded an analysis.")
    document.add_heading("Document coverage", 1)
    coverage = overview["coverage"]
    document.add_paragraph(
        f"Registered files: {coverage['registered']}. Extracted: {coverage['extracted']}. Need attention: {coverage['needs_attention']}. Unsupported: {coverage['unsupported']}. Failed: {coverage['failed']}."
    )
    document.add_paragraph(
        "Extraction and BOQ candidate identification do not establish complete analysis or engineer review. "
        + view["coverage_note"]
    )
    for label, kind in [
        ("Requirements", "requirement"),
        ("Risks", "risk"),
        ("Assumptions", "assumption"),
        ("Questions", "question"),
        ("Observations", "observation"),
    ]:
        selected = [finding for finding in findings if finding["kind"] == kind]
        if not selected:
            continue
        document.add_heading(label, 1)
        for finding in selected:
            document.add_heading(finding["title"], 2)
            document.add_paragraph(
                "State: "
                + finding["state"]
                + ("; source changed" if finding.get("is_stale") else "")
            )
            document.add_paragraph(finding["detail"])
            document.add_paragraph(
                "Evidence: " + (", ".join(finding["source_ids"]) or "No source reference recorded.")
            )
    document.add_heading("Estimate status", 1)
    document.add_paragraph(
        f"BOQ candidates: {len(view['items'])}. Unpriced: {view['unpriced_count']}. Source rows awaiting confirmation: {view['unconfirmed_count']}. VAT treatment unknown: {view['unknown_vat_count']}."
    )
    for total in view["totals"]:
        document.add_paragraph(
            f"{total['currency']} priced subtotal excluding VAT: {total['priced_subtotal_ex_vat']}. Complete total including VAT: {total['total_inc_vat'] or 'not established'}."
        )
    document.add_heading("Source references", 1)
    for source in sources:
        document.add_paragraph(
            f"{source['artifact_name']} — {source['locator']}\nEvidence ID: {source['id']}"
        )
    footer = section.footer.paragraphs[0]
    footer.text = "Draft | "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "Quantix"
    document.core_properties.title = "Tender analysis"
    document.save(path)
