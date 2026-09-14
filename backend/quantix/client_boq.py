"""Price an explicitly mapped copy while preserving the supplied OOXML package."""

import hashlib
import posixpath
import re
from decimal import Decimal, InvalidOperation, localcontext
from io import BytesIO
from pathlib import Path
from typing import Literal
from xml.etree import ElementTree as ET
from xml.parsers import expat
from zipfile import BadZipFile, ZipFile

from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from pydantic import BaseModel, ConfigDict, Field, model_validator

SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
MAX_ORIGINAL_BYTES = 100 * 1024 * 1024
MAX_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_XML_BYTES = 48 * 1024 * 1024
CELL = re.compile(r"^([A-Z]{1,3})([1-9][0-9]{0,6})$")


class ClientBoqModel(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)


class ClientBoqMapping(ClientBoqModel):
    sheet: str = Field(min_length=1, max_length=31)
    source_ids: list[str] = Field(min_length=1, max_length=50)
    item_ids: list[str] = Field(min_length=1, max_length=1000)
    rate_column: str = Field(pattern=r"^[A-Z]{1,3}$")
    amount_column: str = Field(pattern=r"^[A-Z]{1,3}$")
    quantity_column: str | None = Field(default=None, pattern=r"^[A-Z]{1,3}$")

    @model_validator(mode="after")
    def distinct_columns(self):
        columns = [self.rate_column, self.amount_column]
        if self.quantity_column:
            columns.append(self.quantity_column)
        if len(set(columns)) != len(columns):
            raise ValueError("Quantity, rate and amount must use different columns.")
        if any(column_index_from_string(column) > 16384 for column in columns):
            raise ValueError("Use worksheet columns between A and XFD.")
        if len(set(self.item_ids)) != len(self.item_ids):
            raise ValueError("Select each BOQ row once within a mapping.")
        return self


class ClientBoqInput(ClientBoqModel):
    artifact_id: str = Field(min_length=1, max_length=100)
    mappings: list[ClientBoqMapping] = Field(min_length=1, max_length=100)
    currency: str = Field(pattern=r"^[A-Z]{3}$")
    tax_basis: Literal["excluding_vat", "including_vat"]
    mapping_reviewed: Literal[True]
    quantity_mapping_approved: bool = False

    @model_validator(mode="after")
    def explicit_quantity_decision(self):
        if (
            any(mapping.quantity_column for mapping in self.mappings)
            and not self.quantity_mapping_approved
        ):
            raise ValueError(
                "Changing a quantity column requires explicit quantity mapping approval."
            )
        identifiers = [identifier for mapping in self.mappings for identifier in mapping.item_ids]
        if len(set(identifiers)) != len(identifiers):
            raise ValueError("A BOQ row can appear in only one client workbook mapping.")
        if len(identifiers) > 2000:
            raise ValueError("Select at most 2000 rows in one client workbook export.")
        return self


def _xml(payload):
    if len(payload) > MAX_XML_BYTES or re.search(rb"<!DOCTYPE|<!ENTITY", payload, re.I):
        raise ValueError(
            "The workbook XML is too large or contains unsupported entity declarations."
        )
    if payload.startswith((b"\xff\xfe", b"\xfe\xff")) or b"\x00" in payload[:100]:
        raise ValueError(
            "This workbook uses an XML encoding that cannot be edited while preserving its original bytes."
        )
    try:
        return ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ValueError("The original workbook contains unreadable XML.") from exc


def _original(repo, tender_id, artifact_id):
    artifact = repo.get_artifact(tender_id, artifact_id)
    if not artifact["is_current"]:
        raise ValueError("Select the current source workbook revision.")
    if Path(artifact["name"]).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Client-format copies support supplied XLSX and XLSM workbooks only.")
    path = repo.object_path(tender_id, artifact_id)
    if path.stat().st_size > MAX_ORIGINAL_BYTES:
        raise ValueError("This workbook is too large for a client-format copy.")
    payload = path.read_bytes()
    if hashlib.sha256(payload).hexdigest() != artifact["content_hash"]:
        raise ValueError("The preserved workbook changed; import its current original again.")
    return artifact, payload


def _archive(payload):
    try:
        archive = ZipFile(BytesIO(payload))
    except BadZipFile as exc:
        raise ValueError(
            "The workbook is encrypted or is not a readable XLSX/XLSM package. Supply an unencrypted copy."
        ) from exc
    names = archive.namelist()
    if len(names) != len(set(names)) or len(names) > 10000:
        archive.close()
        raise ValueError("The workbook contains duplicate or excessive package entries.")
    if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
        archive.close()
        raise ValueError("The workbook expands beyond the supported package size.")
    if any(info.flag_bits & 1 for info in archive.infolist()):
        archive.close()
        raise ValueError("Encrypted workbook entries cannot be edited.")
    if any(name.lower().startswith("_xmlsignatures/") for name in names):
        archive.close()
        raise ValueError(
            "This workbook has a package digital signature. Provide an unsigned working copy before pricing it."
        )
    return archive


def _read(archive, member):
    try:
        return archive.read(member)
    except (BadZipFile, RuntimeError, NotImplementedError) as exc:
        raise ValueError(
            "An original workbook package entry is damaged, encrypted or uses unsupported compression."
        ) from exc


def _parts(archive):
    try:
        workbook_bytes = _read(archive, "xl/workbook.xml")
        workbook = _xml(workbook_bytes)
        relationships = _xml(_read(archive, "xl/_rels/workbook.xml.rels"))
    except KeyError as exc:
        raise ValueError("The workbook is missing its main sheet relationships.") from exc
    if workbook.tag != f"{{{SHEET_NS}}}workbook":
        raise ValueError("This workbook uses an unsupported spreadsheet XML format.")
    if workbook.find(f"{{{SHEET_NS}}}workbookProtection") is not None:
        raise ValueError(
            "The source workbook is protected. Supply an authorised unprotected working copy."
        )
    links = {}
    for relationship in relationships.findall(f"{{{PACKAGE_NS}}}Relationship"):
        if relationship.get("TargetMode") == "External":
            continue
        target = relationship.get("Target", "")
        member = posixpath.normpath(
            target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
        )
        if member.startswith("../") or member not in archive.namelist():
            continue
        identifier = relationship.get("Id")
        if identifier in links:
            raise ValueError("The workbook contains ambiguous sheet relationship identifiers.")
        links[identifier] = member
    sheets = {}
    for sheet in workbook.findall(f"{{{SHEET_NS}}}sheets/{{{SHEET_NS}}}sheet"):
        member = links.get(sheet.get(f"{{{REL_NS}}}id"))
        if member:
            name = sheet.get("name")
            if name in sheets:
                raise ValueError("The workbook contains ambiguous duplicate worksheet names.")
            sheets[name] = member
    strings = []
    if "xl/sharedStrings.xml" in archive.namelist():
        for item in _xml(_read(archive, "xl/sharedStrings.xml")).findall(f"{{{SHEET_NS}}}si"):
            strings.append("".join(node.text or "" for node in item.iter(f"{{{SHEET_NS}}}t")))
    return workbook_bytes, sheets, strings


def _coordinate(value):
    match = CELL.fullmatch(value or "")
    if not match or int(match[2]) > 1048576 or column_index_from_string(match[1]) > 16384:
        raise ValueError("A source cell has an unsupported worksheet address.")
    return match[1], int(match[2])


def _sheet(archive, member, strings):
    payload = _read(archive, member)
    root = _xml(payload)
    if root.tag != f"{{{SHEET_NS}}}worksheet":
        raise ValueError("The selected sheet is not a standard worksheet.")
    if root.find(f"{{{SHEET_NS}}}sheetProtection") is not None:
        raise ValueError(
            "A selected worksheet is protected. Supply an authorised unprotected working copy."
        )
    cells = {}
    for element in root.findall(f"{{{SHEET_NS}}}sheetData/{{{SHEET_NS}}}row/{{{SHEET_NS}}}c"):
        coordinate = element.get("r")
        _coordinate(coordinate)
        if coordinate in cells:
            raise ValueError("The source worksheet contains duplicate cell addresses.")
        formula = element.find(f"{{{SHEET_NS}}}f")
        value_element = element.find(f"{{{SHEET_NS}}}v")
        value = value_element.text if value_element is not None else None
        kind = element.get("t", "n")
        if kind == "s" and value is not None:
            try:
                index = int(value)
                if index < 0:
                    raise ValueError
                value = strings[index]
            except (ValueError, IndexError) as exc:
                raise ValueError("The workbook has an invalid shared-string reference.") from exc
        elif kind == "inlineStr":
            value = "".join(node.text or "" for node in element.iter(f"{{{SHEET_NS}}}t"))
        cells[coordinate] = {
            "value": value,
            "kind": kind,
            "formula": formula.text if formula is not None else None,
            "formula_attributes": dict(formula.attrib) if formula is not None else {},
        }
    merged = [
        range_boundaries(item.get("ref"))
        for item in root.findall(f"{{{SHEET_NS}}}mergeCells/{{{SHEET_NS}}}mergeCell")
    ]
    return {"payload": payload, "cells": cells, "merged": merged}


def _number(value, description):
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(description + " must have an established numeric value.") from exc
    if not number.is_finite() or number < 0:
        raise ValueError(description + " must be a finite nonnegative number.")
    return number


def _header_sources(repo, tender_id, artifact_id, mapping, sheet, first_row):
    columns = {mapping.rate_column, mapping.amount_column}
    if mapping.quantity_column:
        columns.add(mapping.quantity_column)
    found = set()
    for identifier in mapping.source_ids:
        source = repo.get_evidence(tender_id, identifier)
        if source["artifact_id"] != artifact_id or source.get("sheet") != mapping.sheet:
            raise ValueError(
                "Column mapping evidence must come from the selected workbook and sheet."
            )
        for cell in source.get("metadata", {}).get("cells", []):
            coordinate = cell.get("coordinate", "")
            column, row = _coordinate(coordinate)
            actual = sheet["cells"].get(coordinate, {})
            text = cell.get("value")
            if (
                column in columns
                and row < first_row
                and isinstance(text, str)
                and text.strip()
                and not cell.get("formula")
            ):
                if (
                    str(actual.get("value", "")).strip() != text.strip()
                    or actual.get("formula") is not None
                ):
                    raise ValueError(
                        "The mapped column header differs from its saved source evidence."
                    )
                found.add(column)
    if found != columns:
        raise ValueError(
            "Select source header evidence above the chosen rows for every mapped rate, amount and quantity column."
        )


def _ordinary_amount_formula(formula, quantity_cell, rate_cell):
    value = re.sub(r"\s+", "", formula or "").replace("$", "").upper()
    product = {quantity_cell + "*" + rate_cell, rate_cell + "*" + quantity_cell}
    return value in product or any(value == f"ROUND({expression},2)" for expression in product)


def prepare_client_boq(repo, tender_id, request, estimate):
    """Resolve engineer mappings against exact original cells before making a new copy."""
    request = ClientBoqInput.model_validate(request)
    artifact, payload = _original(repo, tender_id, request.artifact_id)
    if estimate["refresh_required"]:
        raise ValueError(
            "Refresh the BOQ candidates after source changes before making a client-format copy."
        )
    items = {item["id"]: item for item in estimate["items"]}
    changes, references, selected, warnings = [], set(), [], []
    with _archive(payload) as archive:
        _, sheet_members, strings = _parts(archive)
        suffix = (
            ".xlsm"
            if Path(artifact["name"]).suffix.lower() == ".xlsm"
            or any(name.lower().endswith("/vbaproject.bin") for name in archive.namelist())
            else ".xlsx"
        )
        loaded, targets = {}, set()
        for mapping in request.mappings:
            member = sheet_members.get(mapping.sheet)
            if member is None:
                raise ValueError("A mapped worksheet does not exist in the selected original.")
            if member not in loaded:
                loaded[member] = _sheet(archive, member, strings)
            sheet = loaded[member]
            mapped_items = []
            for identifier in mapping.item_ids:
                item = items.get(identifier)
                if (
                    item is None
                    or item["artifact_id"] != artifact["id"]
                    or item["sheet"] != mapping.sheet
                ):
                    raise ValueError(
                        "Select current BOQ rows from the mapped source workbook and sheet."
                    )
                if (
                    not item["confirmed"]
                    or item["unit_rate"] is None
                    or item["rate_ex_vat"] is None
                    or item["line_ex_vat"] is None
                ):
                    raise ValueError(
                        "Every selected row must have confirmed source quantities and an approved price."
                    )
                if item["currency"] != request.currency:
                    raise ValueError(
                        "Every selected row must use the explicitly reviewed workbook currency."
                    )
                if item["tax_basis"] == "unknown" or item["vat_percent"] is None:
                    raise ValueError(
                        "Establish the selected rows' VAT treatment before writing the client workbook."
                    )
                source = repo.get_evidence(tender_id, item["source_id"])
                if source["artifact_id"] != artifact["id"] or source.get("sheet") != mapping.sheet:
                    raise ValueError(
                        "The selected BOQ row source differs from the mapped workbook sheet."
                    )
                row = source.get("metadata", {}).get("row")
                if not isinstance(row, int) or isinstance(row, bool) or not 1 <= row <= 1048576:
                    raise ValueError(
                        "A selected BOQ source does not identify one original worksheet row."
                    )
                quantity_column, quantity_row = _coordinate(item["quantity_cell"])
                if quantity_row != row:
                    raise ValueError("The confirmed quantity cell is outside its BOQ source row.")
                actual_quantity = sheet["cells"].get(item["quantity_cell"], {})
                supplied = _number(item["supplied_quantity"], "The supplied quantity")
                if (
                    actual_quantity.get("kind") not in {"n", None}
                    or _number(actual_quantity.get("value"), "The original quantity") != supplied
                ):
                    raise ValueError(
                        "The original quantity differs from the confirmed saved source quantity."
                    )
                if mapping.quantity_column and mapping.quantity_column != quantity_column:
                    raise ValueError(
                        "Quantity mapping must name the actual confirmed source quantity column."
                    )
                used = _number(item["effective_quantity"], "The quantity used for pricing")
                if used != supplied and not mapping.quantity_column:
                    raise ValueError(
                        "An approved measured quantity differs from the original. Explicitly approve its actual quantity-column mapping or retain the supplied quantity basis."
                    )
                if used != supplied and item["quantity_basis"] != "approved_measurement":
                    raise ValueError("A changed quantity needs a separately approved measurement.")
                unit_column, _ = _coordinate(item["unit_cell"])
                if {mapping.rate_column, mapping.amount_column} & {quantity_column, unit_column}:
                    raise ValueError(
                        "Price columns must not overwrite supplied quantity or unit cells."
                    )
                rate_cell, amount_cell = (
                    mapping.rate_column + str(row),
                    mapping.amount_column + str(row),
                )
                with localcontext() as context:
                    context.prec = 50
                    unit_rate = _number(item["rate_ex_vat"], "The rate excluding VAT")
                    if request.tax_basis == "including_vat":
                        unit_rate = (
                            _number(item["unit_rate"], "The approved rate")
                            if item["tax_basis"] == "including_vat"
                            else unit_rate
                            * (Decimal(1) + _number(item["vat_percent"], "VAT percent") / 100)
                        )
                amount = (
                    item["line_inc_vat"]
                    if request.tax_basis == "including_vat"
                    else item["line_ex_vat"]
                )
                planned = [
                    (rate_cell, format(unit_rate, "f"), "rate"),
                    (amount_cell, format(_number(amount, "The line amount"), "f"), "amount"),
                ]
                if mapping.quantity_column and used != supplied:
                    planned.append((item["quantity_cell"], format(used, "f"), "quantity"))
                for coordinate, value, role in planned:
                    key = (member, coordinate)
                    if key in targets:
                        raise ValueError("Two mappings would write the same original cell.")
                    targets.add(key)
                    column, target_row = _coordinate(coordinate)
                    index = column_index_from_string(column)
                    if any(
                        left <= index <= right and top <= target_row <= bottom
                        for left, top, right, bottom in sheet["merged"]
                    ):
                        raise ValueError(
                            "A mapped price or quantity cell is merged. Resolve the mapping in an authorised working copy."
                        )
                    before = sheet["cells"].get(coordinate, {})
                    formula = before.get("formula")
                    if before.get("formula_attributes") or formula is not None:
                        if (
                            role != "amount"
                            or before.get("formula_attributes")
                            or not _ordinary_amount_formula(
                                formula, item["quantity_cell"], rate_cell
                            )
                        ):
                            raise ValueError(
                                "A mapped target contains a supplied formula that cannot be preserved with the reviewed quantity and rate columns. It has not been overwritten or repaired."
                            )
                        with localcontext() as context:
                            context.prec = 100
                            product = used * unit_rate
                        if not re.sub(r"\s+", "", formula).upper().startswith(
                            "ROUND("
                        ) and product != _number(value, "The approved line amount"):
                            raise ValueError(
                                "A supplied amount formula does not round the mapped rate and quantity to the approved line amount. It has been preserved; resolve the rounding rule in an authorised working copy."
                            )
                    elif (
                        before.get("kind") in {"s", "inlineStr", "str", "e", "b", "d"}
                        and str(before.get("value") or "").strip()
                    ):
                        raise ValueError(
                            "A mapped target contains text, an error or a non-price value. Review the original column mapping."
                        )
                    changes.append(
                        {
                            "sheet": mapping.sheet,
                            "member": member,
                            "cell": coordinate,
                            "role": role,
                            "value": value,
                            "item_id": identifier,
                            "source_id": item["source_id"],
                            "previous_value": before.get("value"),
                            "preserved_formula": formula,
                        }
                    )
                references.update(mapping.source_ids)
                references.add(item["source_id"])
                references.update((item.get("provenance") or {}).get("source_ids", []))
                for proposal in item.get("quantity_proposals", []):
                    if proposal["status"] == "approved":
                        references.update(proposal["source_ids"])
                selected.append(item)
                mapped_items.append(row)
            _header_sources(repo, tender_id, artifact["id"], mapping, sheet, min(mapped_items))
        if any(name.startswith("xl/externalLinks/") for name in archive.namelist()):
            warnings.append("External workbook links are preserved and were not refreshed.")
        if suffix == ".xlsm":
            warnings.append(
                "Supplied VBA and workbook objects are preserved. No macro was executed or validated."
            )
    candidate_count = sum(item["artifact_id"] == artifact["id"] for item in estimate["items"])
    warnings.extend(
        [
            f"Client-format copy prices {len(selected)} selected confirmed rows from {candidate_count} identified source-workbook candidates. Candidate identification does not establish complete BOQ coverage.",
            "Only explicitly mapped cells, any required worksheet range bounds and workbook recalculation settings are changed. Other supplied formulas, cached totals, layout and package objects remain preserved. Open the copy in Excel and review recalculated totals before final approval.",
            "The workbook currency and rate/amount tax basis are the engineer reviewed mapping inputs. Supplied headings and tax notes are preserved and must be checked for agreement.",
        ]
    )
    for identifier in references:
        source = repo.get_evidence(tender_id, identifier)
        if not repo.get_artifact(tender_id, source["artifact_id"])["is_current"]:
            raise ValueError("A selected price, quantity or mapping source is no longer current.")
    return {
        "artifact_id": artifact["id"],
        "source_name": artifact["name"],
        "source_hash": artifact["content_hash"],
        "source_version": artifact["version"],
        "suffix": suffix,
        "changes": changes,
        "source_ids": sorted(references),
        "selected_items": selected,
        "selected_row_count": len(selected),
        "candidate_row_count": candidate_count,
        "pricing_complete": True,
        "currency": request.currency,
        "tax_basis": request.tax_basis,
        "warnings": warnings,
    }


def _tag_end(payload, start):
    quote = None
    for index in range(start, len(payload)):
        value = payload[index]
        if quote:
            if value == quote:
                quote = None
        elif value in {34, 39}:
            quote = value
        elif value == 62:
            return index + 1
    raise ValueError("A workbook XML tag is incomplete.")


def _spans(payload, wanted, row_numbers=None):
    """Locate original byte ranges; serializers must not rewrite unrelated namespaces/objects."""
    _xml(payload)
    parser = expat.ParserCreate(namespace_separator="}")
    stack, spans = [], []

    def start(name, attributes):
        local = name.rsplit("}", 1)[-1]
        position = parser.CurrentByteIndex
        node = {
            "name": local,
            "attributes": attributes,
            "start": position,
            "open_end": _tag_end(payload, position),
            "path": tuple(item["name"] for item in stack) + (local,),
        }
        node["row_number"] = (
            attributes.get("r")
            if local == "row"
            else stack[-1].get("row_number")
            if stack
            else None
        )
        stack.append(node)

    def end(_name):
        node = stack.pop()
        position = parser.CurrentByteIndex
        self_closed = payload[node["start"] : node["open_end"]].rstrip().endswith(b"/>")
        node["close_start"] = node["open_end"] if self_closed else position
        node["end"] = node["open_end"] if self_closed else _tag_end(payload, position)
        if node["path"] in wanted and (
            row_numbers is None or node["row_number"] is None or node["row_number"] in row_numbers
        ):
            spans.append(node)

    parser.StartElementHandler = start
    parser.EndElementHandler = end
    try:
        parser.Parse(payload, True)
    except expat.ExpatError as exc:
        raise ValueError("The workbook XML could not be mapped safely.") from exc
    return spans


def _qname(fragment):
    match = re.match(rb"<([A-Za-z_][A-Za-z0-9_.:-]*)\b", fragment)
    if not match:
        raise ValueError("An original workbook element has an unsupported name.")
    return match[1]


def _numeric_cell(fragment, value, children):
    name = _qname(fragment)
    prefix = name[:-1] if name.endswith(b"c") else b""
    opening_end = _tag_end(fragment, 0)
    opening = fragment[:opening_end]
    opening = re.sub(rb"\s+t\s*=\s*(?:\"[^\"]*\"|'[^']*')", b"", opening)
    numeric = b"<" + prefix + b"v>" + value.encode("ascii") + b"</" + prefix + b"v>"
    if opening.rstrip().endswith(b"/>"):
        return opening.rstrip()[:-2] + b">" + numeric + b"</" + name + b">"
    modifications = [
        (node["start"], node["end"], b"") for node in children if node["name"] in {"v", "is"}
    ]
    insertion = min((node["start"] for node in children), default=fragment.rfind(b"</"))
    if insertion < 0:
        raise ValueError("An original cell has no safe position for its numeric value.")
    modifications.append((insertion, insertion, numeric))
    result = fragment
    for start, end, replacement in sorted(modifications, reverse=True):
        result = result[:start] + replacement + result[end:]
    # Preserve the exact original formula, style and cell attributes except its numeric type.
    return opening + result[opening_end:]


def _patch_sheet(payload, changes):
    wanted = {
        ("worksheet", "sheetData", "row"),
        ("worksheet", "sheetData", "row", "c"),
        ("worksheet", "sheetData", "row", "c", "v"),
        ("worksheet", "sheetData", "row", "c", "is"),
        ("worksheet", "sheetData", "row", "c", "extLst"),
        ("worksheet", "dimension"),
    }
    spans = _spans(payload, wanted, {str(_coordinate(change["cell"])[1]) for change in changes})
    cells = {node["attributes"].get("r"): node for node in spans if node["name"] == "c"}
    rows = {node["attributes"].get("r"): node for node in spans if node["name"] == "row"}
    modifications, additions = [], {}
    for change in changes:
        coordinate = change["cell"]
        cell = cells.get(coordinate)
        if cell:
            fragment = payload[cell["start"] : cell["end"]]
            children = [
                {
                    "name": node["name"],
                    "start": node["start"] - cell["start"],
                    "end": node["end"] - cell["start"],
                }
                for node in spans
                if node["name"] in {"v", "is", "extLst"}
                and cell["open_end"] <= node["start"] < cell["end"]
            ]
            replacement = _numeric_cell(fragment, change["value"], children)
            modifications.append((cell["start"], cell["end"], replacement))
        else:
            column, row_number = _coordinate(coordinate)
            row = rows.get(str(row_number))
            if not row:
                raise ValueError("A mapped original worksheet row is missing.")
            following = [
                node
                for address, node in cells.items()
                if _coordinate(address)[1] == row_number
                and column_index_from_string(_coordinate(address)[0])
                > column_index_from_string(column)
            ]
            position = min((node["start"] for node in following), default=row["close_start"])
            if position < row["open_end"]:
                raise ValueError("A mapped worksheet row cannot accept a new price cell safely.")
            row_name = _qname(payload[row["start"] : row["open_end"]])
            prefix = row_name[:-3] if row_name.endswith(b"row") else b""
            cell_xml = (
                b"<"
                + prefix
                + b'c r="'
                + coordinate.encode("ascii")
                + b'"><'
                + prefix
                + b"v>"
                + change["value"].encode("ascii")
                + b"</"
                + prefix
                + b"v></"
                + prefix
                + b"c>"
            )
            additions.setdefault(position, []).append((column_index_from_string(column), cell_xml))
    modifications.extend(
        (position, position, b"".join(fragment for _, fragment in sorted(values)))
        for position, values in additions.items()
    )
    dimensions = [node for node in spans if node["name"] == "dimension"]
    if dimensions:
        dimension = dimensions[0]
        try:
            left, top, right, bottom = range_boundaries(dimension["attributes"].get("ref", "A1"))
        except ValueError as exc:
            raise ValueError("The original worksheet dimension is invalid.") from exc
        for change in changes:
            column, row = _coordinate(change["cell"])
            left, right = (
                min(left, column_index_from_string(column)),
                max(right, column_index_from_string(column)),
            )
            top, bottom = min(top, row), max(bottom, row)
        updated_bounds = (left, top, right, bottom)
        if updated_bounds != range_boundaries(dimension["attributes"].get("ref", "A1")):
            reference = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}".encode(
                "ascii"
            )
            fragment = payload[dimension["start"] : dimension["end"]]
            fragment = re.sub(
                rb"\bref\s*=\s*(?:\"[^\"]*\"|'[^']*')", b'ref="' + reference + b'"', fragment
            )
            modifications.append((dimension["start"], dimension["end"], fragment))
    for start, end, replacement in sorted(modifications, reverse=True):
        payload = payload[:start] + replacement + payload[end:]
    return payload


def _recalculation(payload):
    root = _xml(payload)
    wanted = {("workbook",)} | {("workbook", child.tag.rsplit("}", 1)[-1]) for child in root}
    spans = _spans(payload, wanted)
    existing = next((node for node in spans if node["name"] == "calcPr"), None)
    if existing:
        opening = payload[existing["start"] : existing["open_end"]]
        for attribute in (b"calcMode", b"fullCalcOnLoad", b"forceFullCalc"):
            opening = re.sub(rb"\s+" + attribute + rb"\s*=\s*(?:\"[^\"]*\"|'[^']*')", b"", opening)
        insertion = (
            opening.rfind(b"/>") if opening.rstrip().endswith(b"/>") else opening.rfind(b">")
        )
        opening = (
            opening[:insertion]
            + b' calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"'
            + opening[insertion:]
        )
        return payload[: existing["start"]] + opening + payload[existing["open_end"] :]
    root_span = next(node for node in spans if node["path"] == ("workbook",))
    root_name = _qname(payload[root_span["start"] : root_span["open_end"]])
    prefix = root_name[:-8] if root_name.endswith(b"workbook") else b""
    later = {
        "oleSize",
        "customWorkbookViews",
        "pivotCaches",
        "smartTagPr",
        "smartTagTypes",
        "webPublishing",
        "fileRecoveryPr",
        "webPublishObjects",
        "extLst",
    }
    position = min(
        (node["start"] for node in spans if node["name"] in later), default=root_span["close_start"]
    )
    if position < 0:
        raise ValueError("The workbook has no safe location for recalculation settings.")
    addition = b"<" + prefix + b'calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
    return payload[:position] + addition + payload[position:]


def write_client_boq(path, repo, tender_id, prepared):
    artifact, payload = _original(repo, tender_id, prepared["artifact_id"])
    if artifact["content_hash"] != prepared["source_hash"]:
        raise ValueError("The workbook source changed while preparing its client-format copy.")
    by_member = {}
    for change in prepared["changes"]:
        by_member.setdefault(change["member"], []).append(change)
    with _archive(payload) as original, ZipFile(path, "x") as output:
        output.comment = original.comment
        for member in original.infolist():
            content = _read(original, member.filename)
            if member.filename in by_member:
                content = _patch_sheet(content, by_member[member.filename])
            elif member.filename == "xl/workbook.xml":
                content = _recalculation(content)
            # Copy each original entry's metadata and bytes; never load or resave VBA or drawings.
            output.writestr(member, content)
