"""The optional Word boundary is controlled without launching Office in unit tests."""

from __future__ import annotations

import importlib
import importlib.util
import subprocess
import sys
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Windows-only Office boundary")


def worker():
    assert importlib.util.find_spec("quantix.document_word"), "Isolated Word conversion is missing"
    return importlib.import_module("quantix.document_word")


def test_word_conversion_uses_a_copy_and_returns_only_derived_output(tmp_path, monkeypatch):
    module = worker()
    source = tmp_path / "original.doc"
    source.write_bytes(b"source bytes")
    actual_popen = subprocess.Popen
    intermediate = []

    def substitute(command, **kwargs):
        copied, output, state = map(Path, command[-3:])
        assert copied != source and copied.read_bytes() == b"source bytes"
        assert copied.parent == output.parent == state.parent
        intermediate.append(copied.parent)
        return actual_popen(
            [
                sys.executable,
                "-c",
                "import pathlib,sys; pathlib.Path(sys.argv[1]).write_bytes(b'PK converted bytes')",
                str(output),
            ],
            **kwargs,
        )

    monkeypatch.setattr(module.subprocess, "Popen", substitute)
    assert module.convert_legacy_word(source) == b"PK converted bytes"
    assert source.read_bytes() == b"source bytes"
    assert list(tmp_path.iterdir()) == [source]
    assert all(not p.exists() for p in intermediate)


@pytest.mark.parametrize("cancel", [False, True])
def test_word_worker_timeout_and_cancellation_are_bounded(tmp_path, monkeypatch, cancel):
    module = worker()
    source = tmp_path / "original.doc"
    source.write_bytes(b"source bytes")
    actual_popen = subprocess.Popen
    processes = []

    def substitute(command, **kwargs):
        if command[-2:] == ["/T", "/F"]:
            return actual_popen(command, **kwargs)
        process = actual_popen([sys.executable, "-c", "import time; time.sleep(30)"], **kwargs)
        processes.append(process)
        return process

    monkeypatch.setattr(module.subprocess, "Popen", substitute)
    monkeypatch.setattr(module, "WORD_TIMEOUT_SECONDS", 0.4)
    started = time.monotonic()
    with pytest.raises(InterruptedError if cancel else module.WordConversionUnavailable):
        module.convert_legacy_word(
            source, (lambda: time.monotonic() - started > 0.15) if cancel else None
        )
    assert time.monotonic() - started < 3
    assert processes and all(p.poll() is not None for p in processes)
    assert source.read_bytes() == b"source bytes"


def test_word_open_disables_macros_prompts_links_and_preserves_settings(tmp_path):
    module = worker()
    source, output = tmp_path / "input.doc", tmp_path / "output.docx"
    source.write_bytes(b"source")
    closed = []
    app = SimpleNamespace(
        AutomationSecurity=1,
        DisplayAlerts=-1,
        Visible=True,
        Options=SimpleNamespace(UpdateLinksAtOpen=True),
    )

    class Document:
        def SaveAs2(self, **kwargs):
            assert kwargs == {"FileName": str(output), "FileFormat": 16, "AddToRecentFiles": False}
            output.write_bytes(b"derived only")

        def Close(self, **kwargs):
            assert kwargs == {"SaveChanges": 0}
            closed.append(True)

    def open_document(**kwargs):
        # This boundary rejects unsafe opens; it does not merely observe calls.
        assert app.AutomationSecurity == 3
        assert app.DisplayAlerts == 0 and app.Visible is False
        assert app.Options.UpdateLinksAtOpen is False
        assert kwargs["FileName"] == str(source)
        assert kwargs["ReadOnly"] is True and kwargs["AddToRecentFiles"] is False
        assert kwargs["Visible"] is False and kwargs["ConfirmConversions"] is False
        assert kwargs["NoEncodingDialog"] is True and kwargs["OpenAndRepair"] is False
        return Document()

    app.Documents = SimpleNamespace(Open=open_document)
    module._save_document(app, source, output)
    assert output.read_bytes() == b"derived only" and source.read_bytes() == b"source"
    assert closed and app.AutomationSecurity == 1 and app.DisplayAlerts == -1
    assert app.Options.UpdateLinksAtOpen is True


def test_word_failure_is_an_explicit_coverage_exception(tmp_path, monkeypatch):
    module = worker()
    documents = importlib.import_module("quantix.documents")
    source = tmp_path / "binary.doc"
    source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 40)

    def unavailable(*args, **kwargs):
        raise module.WordConversionUnavailable("Word conversion is unavailable")

    monkeypatch.setattr(module, "convert_legacy_word", unavailable)
    result = documents.extract_document(source, source.name)
    assert result.kind == "word" and result.status == "unsupported"
    assert any(w["code"] == "doc_conversion_unavailable" for w in result.warnings)


def test_worker_owns_only_a_new_word_process_without_application_hwnd(tmp_path, monkeypatch):
    module = worker()
    import win32api
    import win32com.client
    import win32process

    closes = []
    app = SimpleNamespace(Quit=lambda **kwargs: closes.append(kwargs))

    # PyHANDLE is deliberately not a context manager; Word.Application has no Hwnd.
    class Handle:
        def Close(self):
            pass

    processes = iter([[100], [100, 200]])
    monkeypatch.setattr(win32process, "EnumProcesses", lambda: next(processes))
    monkeypatch.setattr(win32com.client, "DispatchEx", lambda name: app)
    monkeypatch.setattr(win32api, "OpenProcess", lambda access, inherit, pid: Handle())
    monkeypatch.setattr(
        win32process,
        "GetModuleFileNameEx",
        lambda handle, mod: r"C:\Program Files\Microsoft Office\WINWORD.EXE",
    )
    monkeypatch.setattr(
        win32process,
        "GetProcessTimes",
        lambda handle: {"CreationTime": datetime(2026, 1, 1, tzinfo=timezone.utc)},
    )
    monkeypatch.setattr(
        module,
        "_save_document",
        lambda application, source, target: target.write_bytes(b"converted"),
    )
    source, output, state = (
        tmp_path / "input.doc",
        tmp_path / "output.docx",
        tmp_path / "state.json",
    )
    assert module._run(source, output, state) == 0
    assert output.read_bytes() == b"converted"
    assert '"pid": 200' in state.read_text()
    assert closes == [{"SaveChanges": 0}]


def test_word_conversions_wait_cancellably_for_exclusive_office_access(tmp_path):
    module = worker()
    source = tmp_path / "source.doc"
    source.write_bytes(b"test")
    started = threading.Event()
    stop = threading.Event()
    errors = []

    def convert():
        started.set()
        try:
            module.convert_legacy_word(source, stop.is_set)
        except InterruptedError:
            errors.append("cancelled")

    assert hasattr(module, "WORD_LOCK"), (
        "Concurrent conversions must not confuse Word process ownership"
    )
    with module.WORD_LOCK:
        thread = threading.Thread(target=convert)
        thread.start()
        started.wait(1)
        stop.set()
        thread.join(1)
        assert not thread.is_alive()
    assert errors == ["cancelled"]


def test_successful_legacy_conversion_retains_source_format_and_locator_basis(
    tmp_path, monkeypatch
):
    module = worker()
    from docx import Document

    documents = importlib.import_module("quantix.documents")
    source = tmp_path / "binary.doc"
    source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 40)
    derived = tmp_path / "fixture.docx"
    doc = Document()
    doc.add_paragraph("Original requirement")
    doc.save(derived)
    monkeypatch.setattr(module, "convert_legacy_word", lambda *args: derived.read_bytes())
    result = documents.extract_document(source, source.name)
    assert result.status == "extracted"
    assert result.segments[0].text == "Original requirement"
    assert result.metadata["format"] == "ole_word"
    assert result.metadata["source_locator_basis"] == "converted_docx_structure"


def test_settings_restore_even_if_word_document_close_fails(tmp_path):
    module = worker()
    app = SimpleNamespace(
        AutomationSecurity=1,
        DisplayAlerts=-1,
        Visible=True,
        Options=SimpleNamespace(UpdateLinksAtOpen=True),
    )

    class Document:
        def SaveAs2(self, **kwargs):
            raise OSError("Unable to save conversion")

        def Close(self, **kwargs):
            raise OSError("Unable to close document")

    app.Documents = SimpleNamespace(Open=lambda **kwargs: Document())
    with pytest.raises(OSError):
        module._save_document(app, tmp_path / "a.doc", tmp_path / "b.docx")
    assert app.AutomationSecurity == 1 and app.DisplayAlerts == -1
    assert app.Options.UpdateLinksAtOpen is True
