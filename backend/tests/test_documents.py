"""Document ingestion preserves evidence; fixtures never use customer documents."""

from __future__ import annotations

import hashlib
import importlib
import importlib.util
import io
import json
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image


def worker():
    assert importlib.util.find_spec("quantix.documents"), "Document extraction is not implemented"
    return importlib.import_module("quantix.documents")


def pdf_bytes(texts: list[str | None]) -> bytes:
    """Small valid PDF, with independent text objects and no third-party writer."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    page_ids = []
    for text in texts:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        stream = (f"BT /F1 12 Tf 30 80 Td ({text}) Tj ET" if text else "").encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 100] /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    kids = " ".join(f"{n} 0 R" for n in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{n} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(data)


def rewrite_zip(path: Path, transform):
    with zipfile.ZipFile(path) as source:
        entries = {n: source.read(n) for n in source.namelist()}
    transform(entries)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries.items():
            target.writestr(name, data)


def workbook_fixture(path: Path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Civil"
    sheet.append(["Description", "Quantity", "Unit", "Amount"])
    # Deliberate header/data mismatch: extraction must not reinterpret columns.
    sheet.append(["Concrete", "m3", 12.5, "=C2*100"])
    sheet.append(["Broken", "m2", 2, "=#REF!*2"])
    sheet.append(["Uncached", None, None, "=SUM(C2:C3)"])
    sheet.append(["Error", None, None, "#DIV/0!"])
    sheet.append(["Dated", date(2026, 8, 1), True, 0])
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["C"].hidden = True
    sheet["D2"].number_format = "#,##0.00"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["B2"] = "Scope note"
    hidden.merge_cells("B2:C2")
    wb.save(path)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

    def patch(entries):
        root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
        for cell in root.findall(".//s:c", ns):
            if cell.attrib["r"] == "D2":
                cell.find("s:v", ns).text = "1250"
            if cell.attrib["r"] == "D3":
                cell.set("t", "e")
                cell.find("s:v", ns).text = "#REF!"
        entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        entries["xl/vbaProject.bin"] = b"SYNTHETIC VBA PRESENCE ONLY; NOT EXECUTABLE"

    rewrite_zip(path, patch)


def test_pdf_extracts_page_sources_and_reports_empty_page(tmp_path):
    module = worker()
    path = tmp_path / "blob"
    path.write_bytes(pdf_bytes(["Concrete grade C30", None, "Second requirement"]))
    original = hashlib.sha256(path.read_bytes()).digest()
    result = module.extract_document(path, "spec.pdf")
    assert result.kind == "pdf"
    assert result.status == "needs_attention"
    assert [(s.page, s.locator, s.text.strip()) for s in result.segments] == [
        (1, "page:1", "Concrete grade C30"),
        (3, "page:3", "Second requirement"),
    ]
    assert result.metadata["page_count"] == 3
    assert any(w["code"] == "pdf_no_text" and w["locator"] == "page:2" for w in result.warnings)
    assert hashlib.sha256(path.read_bytes()).digest() == original


def test_render_returns_real_bounded_png_and_checks_one_based_page(tmp_path):
    module = worker()
    path = tmp_path / "drawing.pdf"
    path.write_bytes(pdf_bytes(["Drawing"]))
    before = path.read_bytes()
    rendered = module.render_pdf_page(path, 1, max_width=400)
    with Image.open(io.BytesIO(rendered)) as image:
        assert image.format == "PNG"
        assert image.size == (400, 200)
    for page in [0, -1, 2]:
        with pytest.raises(ValueError):
            module.render_pdf_page(path, page)
    assert path.read_bytes() == before


def test_spreadsheet_preserves_formulas_caches_errors_and_locations(tmp_path):
    module = worker()
    path = tmp_path / "boq.xlsm"
    workbook_fixture(path)
    before = path.read_bytes()
    result = module.extract_document(path, path.name)
    assert result.kind == "spreadsheet" and result.status == "needs_attention"
    assert result.metadata["has_vba"] is True
    row = next(s for s in result.segments if s.sheet == "Civil" and s.cell_range == "A2:D2")
    cells = {c["coordinate"]: c for c in row.metadata["cells"]}
    assert cells["B2"]["value"] == "m3"
    assert cells["C2"]["value"] == 12.5
    assert cells["C2"]["column_hidden"] is True
    assert cells["D2"]["formula"] == "=C2*100"
    assert cells["D2"]["value"] == "=C2*100"
    assert cells["D2"]["cached_value"] == 1250
    assert cells["D2"]["number_format"] == "#,##0.00"
    assert row.metadata["row_hidden"] is True
    assert row.metadata["column_interpretation"] == "unverified"
    hidden = next(s for s in result.segments if s.sheet == "Hidden")
    assert hidden.metadata["sheet_state"] == "hidden"
    assert "B2:C2" in hidden.metadata["merged_ranges"]
    codes = {w["code"] for w in result.warnings}
    assert {"spreadsheet_error", "formula_cache_missing"} <= codes
    json.dumps([s.metadata for s in result.segments], allow_nan=False)
    assert path.read_bytes() == before


def test_docx_preserves_paragraph_table_order_nested_cells_and_headers(tmp_path):
    module = worker()
    path = tmp_path / "vendors.docx"
    doc = Document()
    doc.add_paragraph("Approved vendors")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cement"
    table.cell(0, 1).text = "Example Supplier"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Conditional approval"
    doc.add_paragraph("Check delivery")
    doc.sections[0].header.paragraphs[0].text = "Tender Rev 3"
    doc.save(path)
    result = module.extract_document(path, path.name)
    assert result.kind == "word" and result.status == "extracted"
    text = "\n".join(s.text for s in result.segments)
    assert text.index("Approved vendors") < text.index("Cement") < text.index("Check delivery")
    assert "Conditional approval" in text and "Tender Rev 3" in text
    source = next(s for s in result.segments if s.text == "Example Supplier")
    assert source.locator == "table:1/row:1/cell:2/paragraph:1"
    assert source.page is None


def test_formats_are_identified_by_content_and_parse_failures_are_explicit(tmp_path):
    module = worker()
    path = tmp_path / "source"
    path.write_bytes(b"AC1032" + b"\0" * 32)
    result = module.extract_document(path, "drawing.dwg")
    assert result.kind == "cad" and result.status == "unsupported"
    assert result.metadata["dwg_version"] == "AC1032"
    path.write_bytes(b"not a pdf")
    result = module.extract_document(path, "drawing.pdf")
    assert result.kind == "pdf" and result.status == "failed"
    path.write_bytes(b"not a supported format")
    assert module.extract_document(path, "unknown.bin").status == "unsupported"


def test_cancellation_raises_before_work_and_midway_through_pdf(tmp_path):
    module = worker()
    path = tmp_path / "spec.pdf"
    path.write_bytes(pdf_bytes(["one", "two", "three"]))
    with pytest.raises(InterruptedError):
        module.extract_document(path, path.name, lambda: True)
    calls = 0

    def cancelled():
        nonlocal calls
        calls += 1
        return calls >= 4

    with pytest.raises(InterruptedError):
        module.extract_document(path, path.name, cancelled)


def test_limits_report_partial_coverage_without_silent_loss(tmp_path, monkeypatch):
    module = worker()
    path = tmp_path / "spec.pdf"
    path.write_bytes(pdf_bytes(["one", "two", "three"]))
    monkeypatch.setattr(module, "MAX_PDF_PAGES", 2)
    result = module.extract_document(path, path.name)
    assert [s.page for s in result.segments] == [1, 2]
    assert result.status == "needs_attention"
    assert any(w["code"] == "page_limit" for w in result.warnings)
    monkeypatch.setattr(module, "MAX_FILE_BYTES", 10)
    result = module.extract_document(path, path.name)
    assert result.status == "needs_attention" and not result.segments
    assert any(w["code"] == "file_limit" for w in result.warnings)


def test_render_and_extract_share_pdfium_lock(tmp_path, monkeypatch):
    module = worker()
    path = tmp_path / "spec.pdf"
    path.write_bytes(pdf_bytes(["one"]))
    real = module.pdfium.PdfDocument
    entered = 0
    guard = threading.Lock()

    def checked(*args, **kwargs):
        nonlocal entered
        assert module.PDFIUM_LOCK.locked(), "PDFium entered outside the shared lock"
        with guard:
            entered += 1
        return real(*args, **kwargs)

    monkeypatch.setattr(module.pdfium, "PdfDocument", checked)
    with ThreadPoolExecutor(max_workers=2) as pool:
        extraction = pool.submit(module.extract_document, path, path.name)
        render = pool.submit(module.render_pdf_page, path, 1)
        assert extraction.result().segments[0].text.strip() == "one"
        assert render.result().startswith(b"\x89PNG")
    assert entered == 2


def test_output_and_row_limits_keep_complete_source_values(tmp_path, monkeypatch):
    module = worker()
    path = tmp_path / "boq.xlsx"
    wb = Workbook()
    for n in range(1, 5):
        wb.active.append([f"Item {n}", n])
    wb.save(path)
    monkeypatch.setattr(module, "MAX_ROWS", 2)
    result = module.extract_document(path, path.name)
    assert len(result.segments) == 2 and result.status == "needs_attention"
    assert any(w["code"] == "row_limit" for w in result.warnings)
    pdf = tmp_path / "spec.pdf"
    pdf.write_bytes(pdf_bytes(["Keep this whole value", "Omit this whole value"]))
    monkeypatch.setattr(module, "MAX_OUTPUT_BYTES", 150)
    result = module.extract_document(pdf, pdf.name)
    assert len(result.segments) == 1
    assert result.segments[0].text.strip() == "Keep this whole value"
    assert any(w["code"] == "output_limit" for w in result.warnings)


def test_office_zip_limits_and_xml_entity_declarations_are_rejected(tmp_path, monkeypatch):
    module = worker()
    path = tmp_path / "bad.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", b'<!DOCTYPE x [<!ENTITY q "expanded">]><x>&q;</x>')
    result = module.extract_document(path, path.name)
    assert result.status == "failed" and not result.segments
    assert any(w["code"] == "parse_error" for w in result.warnings)
    monkeypatch.setattr(module, "MAX_ZIP_BYTES", 10)
    result = module.extract_document(path, path.name)
    assert result.status == "needs_attention"
    assert any(w["code"] == "package_limit" for w in result.warnings)


def test_docx_unread_revisions_are_visible_and_rtf_is_not_called_binary_doc(tmp_path):
    module = worker()
    path = tmp_path / "revision.docx"
    doc = Document()
    doc.add_paragraph("Current scope")
    doc.save(path)

    def patch(entries):
        data = entries["word/document.xml"]
        entries["word/document.xml"] = data.replace(
            b"</w:body>", b"<w:ins><w:p><w:r><w:t>Proposed scope</w:t></w:r></w:p></w:ins></w:body>"
        )

    rewrite_zip(path, patch)
    result = module.extract_document(path, path.name)
    assert result.status == "needs_attention"
    assert any(w["code"] == "word_revisions" for w in result.warnings)
    path.write_bytes(b"{\\rtf1 actual rtf}")
    result = module.extract_document(path, "renamed.doc")
    assert result.metadata["format"] == "rtf"


def test_spreadsheet_incorrect_dimension_does_not_hide_real_source_rows(tmp_path):
    module = worker()
    path = tmp_path / "dimensions.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Header"
    wb.active["D5"] = "Important scope"
    wb.save(path)

    def patch(entries):
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
            b'ref="A1:D5"', b'ref="A1:A1"'
        )

    rewrite_zip(path, patch)
    result = module.extract_document(path, path.name)
    assert any(
        c["coordinate"] == "D5" and c["value"] == "Important scope"
        for s in result.segments
        for c in s.metadata["cells"]
    )
    assert result.metadata["sheets"][0]["max_row"] == 5
    assert result.metadata["sheets"][0]["max_column"] == 4


def test_empty_workbook_and_invalid_zero_page_pdf_do_not_claim_extracted_evidence(tmp_path):
    module = worker()
    workbook = tmp_path / "empty.xlsx"
    Workbook().save(workbook)
    result = module.extract_document(workbook, workbook.name)
    assert result.status == "needs_attention" and not result.segments
    assert any(w["code"] == "spreadsheet_no_values" for w in result.warnings)
    pdf = tmp_path / "empty.pdf"
    pdf.write_bytes(pdf_bytes([]))
    result = module.extract_document(pdf, pdf.name)
    assert result.status == "failed" and not result.segments
    assert any(w["code"] == "parse_error" for w in result.warnings)


def test_spreadsheet_parser_coverage_warning_is_attached_to_document(tmp_path):
    module = worker()
    path = tmp_path / "header.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Scope"
    wb.save(path)

    def patch(entries):
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
            b"</worksheet>",
            b"<headerFooter><oddHeader>Unreadable header control</oddHeader></headerFooter></worksheet>",
        )

    rewrite_zip(path, patch)
    result = module.extract_document(path, path.name)
    assert result.segments[0].metadata["cells"][0]["value"] == "Scope"
    assert result.status == "needs_attention"
    assert any(w["code"] == "spreadsheet_parser_warning" for w in result.warnings)
