"""Real draft artifacts and scope-bound local submission approvals."""

import hashlib
import importlib
import json
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from test_estimates import approval, rate, seed

from quantix.estimates import EstimateService
from quantix.outputs import OutputService
from quantix.repository import Repository


@pytest.fixture
def workspace(tmp_path):
    repo = Repository(tmp_path)
    tid = repo.create_tender("Foundation works")["id"]
    original = b"Synthetic preserved source"
    digest = hashlib.sha256(original).hexdigest()
    seed(repo, tid, digest=digest)
    (repo.objects / digest).write_bytes(original)
    estimates = EstimateService(repo)
    estimates.refresh(tid)
    return repo, tid, estimates, OutputService(repo)


def programme(**updates):
    return {
        "title": "Foundation construction programme",
        "start_date": "2026-09-10",
        "working_week": [0, 1, 2, 3, 4],
        "holidays": ["2026-09-14"],
        "activities": [
            {"id": "exc", "title": "Excavation", "duration_days": 2},
            {"id": "bl", "title": "Blinding", "duration_days": 1, "predecessor_ids": ["exc"]},
            {"id": "re", "title": "Reinforcement", "duration_days": 3, "predecessor_ids": ["exc"]},
            {"id": "con", "title": "Concrete", "duration_days": 1, "predecessor_ids": ["bl", "re"]},
        ],
        "assumptions": ["One working shift per day"],
        **updates,
    }


def completed_task(repo, tid, source):
    plan = repo.create_plan(
        tid,
        "Foundation work",
        [
            {
                "title": "Concrete method",
                "description": "Describe the saved method",
                "role": "Civil engineer",
                "source_ids": [source],
            }
        ],
    )
    repo.approve_plan(tid, plan["id"], "Review the foundation scope")
    task = repo.list_tasks(tid)[0]
    repo.update_task(
        tid,
        task["id"],
        status="completed",
        result={
            "summary": "Inspect formation before blinding concrete.",
            "source_ids": [source],
            "findings": [
                {
                    "title": "Bearing level",
                    "detail": "Confirm founding level before excavation.",
                    "kind": "question",
                    "state": "proposed",
                    "source_ids": [source],
                }
            ],
        },
    )
    return task


def test_technical_document_uses_selected_saved_task_and_provenance(workspace):
    repo, tid, estimates, outputs = workspace
    source = estimates.view(tid)["items"][0]["source_id"]
    task = completed_task(repo, tid, source)
    repo.add_message(tid, "manager", "Unrelated later office summary", [])
    record = outputs.generate(tid, approval(kind="technical_docx", task_id=task["id"]))
    doc = Document(outputs.path(tid, record["id"]))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Inspect formation before blinding concrete." in text
    assert "Unrelated later office summary" not in text
    assert "Confirm founding level" in text
    assert source in text and "Draft" in text and "Source references" in text
    assert record["metadata"]["task_id"] == task["id"]


def test_programme_dates_observe_working_calendar_and_all_predecessors(workspace):
    repo, tid, estimates, outputs = workspace
    record = outputs.generate(tid, approval(kind="programme_xlsx", programme=programme()))
    scheduled = record["metadata"]["scheduled_activities"]
    dates = {r["id"]: (r["start_date"], r["finish_date"]) for r in scheduled}
    assert dates == {
        "exc": ("2026-09-10", "2026-09-11"),
        "bl": ("2026-09-15", "2026-09-15"),
        "re": ("2026-09-15", "2026-09-17"),
        "con": ("2026-09-18", "2026-09-18"),
    }
    book = load_workbook(outputs.path(tid, record["id"]))
    assert "Programme" in book.sheetnames and "Calendar" in book.sheetnames
    assert "One working shift" in str(list(book["Calendar"].values))


@pytest.mark.parametrize(
    "changes",
    [
        {"activities": [{"id": "a", "title": "Invalid", "duration_days": 0}]},
        {
            "activities": [
                {"id": "a", "title": "Invalid", "duration_days": 1, "predecessor_ids": ["missing"]}
            ]
        },
        {
            "activities": [
                {"id": "a", "title": "Invalid", "duration_days": 1, "predecessor_ids": ["a"]}
            ]
        },
        {"working_week": []},
    ],
)
def test_programme_rejects_invalid_calendar_and_dependencies(workspace, changes):
    _, tid, _, outputs = workspace
    with pytest.raises(ValueError):
        outputs.generate(tid, approval(kind="programme_xlsx", programme=programme(**changes)))
    assert outputs.list(tid) == []


def test_register_preserves_decisions_and_inert_imported_text(workspace):
    repo, tid, estimates, outputs = workspace
    source = estimates.view(tid)["items"][0]["source_id"]
    repo.add_finding(tid, '=HYPERLINK("bad")', "Check dewatering", "risk", [source])
    record = outputs.generate(tid, approval(kind="registers_xlsx"))
    book = load_workbook(outputs.path(tid, record["id"]))
    assert book["Risks"]["B7"].data_type == "s"
    assert source in str(list(book["Risks"].values))
    assert "proposed" in str(list(book["Risks"].values))


def submission_service(repo):
    return importlib.import_module("quantix.submissions").SubmissionService(repo)


def release(service, tid, output):
    preview = service.preview(tid, {"output_ids": [output["id"]]})
    values = approval(
        output_ids=[output["id"]],
        fingerprint=preview["fingerprint"],
        final_review_confirmed=True,
        acknowledged_scope="Construction sequence only; commercial offer excluded",
        acknowledged_gaps=preview["warnings"],
    )
    return service.approve(tid, values)


def test_final_export_freezes_exact_files_scope_and_decision_without_transmission(workspace):
    repo, tid, _, outputs = workspace
    output = outputs.generate(tid, approval(kind="programme_xlsx", programme=programme()))
    service = submission_service(repo)
    result = release(service, tid, output)
    assert result["status"] == "approved_export"
    assert result["acknowledged_scope"].startswith("Construction sequence only")
    with ZipFile(service.path(tid, result["id"])) as archive:
        assert archive.read(output["filename"]) == outputs.path(tid, output["id"]).read_bytes()
        manifest = json.loads(archive.read("submission-manifest.json"))
        assert manifest["outputs"][0]["sha256"] == output["sha256"]
        assert manifest["external_transmission"] is False


@pytest.mark.parametrize("mutation", ["bytes", "manifest", "source", "estimate"])
def test_final_export_rejects_changed_output_or_basis(workspace, mutation):
    repo, tid, estimates, outputs = workspace
    item = estimates.view(tid)["items"][0]
    estimates.update_item(tid, item["id"], rate(vat_percent="14"))
    output = outputs.generate(tid, approval(kind="boq_xlsx"))
    path = outputs.path(tid, output["id"])
    service = submission_service(repo)
    preview = service.preview(tid, {"output_ids": [output["id"]]})
    if mutation == "bytes":
        path.write_bytes(b"tampered")
    elif mutation == "manifest":
        path.with_suffix(".json").write_text("{}", encoding="utf-8")
    elif mutation == "source":
        seed(repo, tid, digest="b" * 64)
    else:
        estimates.update_item(tid, item["id"], rate(vat_percent="14") | {"unit_rate": "30"})
    with pytest.raises(ValueError, match="changed|stale|missing|integrity|current"):
        service.approve(
            tid,
            approval(
                output_ids=[output["id"]],
                fingerprint=preview["fingerprint"],
                final_review_confirmed=True,
                acknowledged_scope="Selected BOQ",
                acknowledged_gaps=preview["warnings"],
            ),
        )
    assert service.list(tid) == []


def test_draft_can_have_gaps_but_final_boq_requires_approved_commercial_basis(workspace):
    repo, tid, _, outputs = workspace
    output = outputs.generate(tid, approval(kind="boq_xlsx"))
    service = submission_service(repo)
    assert service.preview(tid, {"output_ids": [output["id"]]})["blocking_reasons"]
    with pytest.raises(ValueError):
        release(service, tid, output)


def test_final_scope_and_current_gap_acknowledgement_are_mandatory(workspace):
    repo, tid, _, outputs = workspace
    output = outputs.generate(tid, approval(kind="programme_xlsx", programme=programme()))
    service = submission_service(repo)
    preview = service.preview(tid, {"output_ids": [output["id"]]})
    with pytest.raises(ValueError):
        service.approve(
            tid,
            approval(
                output_ids=[output["id"]],
                fingerprint=preview["fingerprint"],
                final_review_confirmed=True,
                acknowledged_scope=" ",
                acknowledged_gaps=preview["warnings"],
            ),
        )
    with pytest.raises(ValueError, match="acknowledge|gaps"):
        service.approve(
            tid,
            approval(
                output_ids=[output["id"]],
                fingerprint=preview["fingerprint"],
                final_review_confirmed=True,
                acknowledged_scope="Sequence only",
                acknowledged_gaps=[],
            ),
        )


def test_comparison_uses_real_saved_reply_and_does_not_create_an_offer(workspace):
    repo, tid, _, outputs = workspace
    from quantix.correspondence import QuoteService

    mail = QuoteService(repo)
    quote = mail.create_draft(
        tid,
        {
            "to": ["supplier@example.com"],
            "subject": "Concrete RFQ",
            "body": "Quote foundation concrete",
            "source_ids": [],
        },
    )
    reply = mail.register_reply(
        tid,
        quote["id"],
        {
            "sender": "supplier@example.com",
            "received_at": "2026-09-06T10:00:00+00:00",
            "text": "EGP 1200 per m3 excluding VAT; pump excluded.",
        },
    )
    output = outputs.generate(tid, approval(kind="comparison_xlsx"))
    book = load_workbook(outputs.path(tid, output["id"]))
    assert "EGP 1200 per m3 excluding VAT; pump excluded." in str(
        list(book["Supplier replies"].values)
    )
    assert reply["source_ids"][0] in str(list(book["Sources"].values))
    assert "No rate proposals" in str(list(book["Rate comparison"].values))


def test_unfinished_or_other_tender_task_cannot_supply_technical_document(workspace):
    repo, tid, estimates, outputs = workspace
    source = estimates.view(tid)["items"][0]["source_id"]
    task = completed_task(repo, tid, source)
    other = repo.create_tender("Other")["id"]
    with pytest.raises(KeyError):
        outputs.generate(other, approval(kind="technical_docx", task_id=task["id"]))
    repo.update_task(tid, task["id"], status="failed")
    with pytest.raises(ValueError, match="completed"):
        outputs.generate(tid, approval(kind="technical_docx", task_id=task["id"]))


def test_technical_current_finding_decision_supersedes_saved_proposed_state(workspace):
    repo, tid, estimates, outputs = workspace
    source = estimates.view(tid)["items"][0]["source_id"]
    task = completed_task(repo, tid, source)
    finding = repo.add_finding(
        tid, "Pump allowance", "Include the engineer allowance", "assumption", [source]
    )
    repo.update_task(
        tid,
        task["id"],
        status="completed",
        result={
            "summary": "Use the approved pump allowance.",
            "source_ids": [source],
            "findings": [finding],
        },
    )
    original = outputs.generate(tid, approval(kind="technical_docx", task_id=task["id"]))
    service = submission_service(repo)
    assert service.preview(tid, {"output_ids": [original["id"]]})["blocking_reasons"]
    repo.decide_finding(tid, finding["id"], "accept", "Engineer checked the allowance")
    reviewed = outputs.generate(tid, approval(kind="technical_docx", task_id=task["id"]))
    assert not service.preview(tid, {"output_ids": [reviewed["id"]]})["blocking_reasons"]
    text = "\n".join(p.text for p in Document(outputs.path(tid, reviewed["id"])).paragraphs)
    assert "accepted" in text


@pytest.mark.parametrize("mutation", ["source_bytes", "source_missing", "output_missing"])
def test_final_export_rechecks_preserved_sources_and_file_presence(workspace, mutation):
    repo, tid, estimates, outputs = workspace
    source = estimates.view(tid)["items"][0]["source_id"]
    task = completed_task(repo, tid, source)
    output = outputs.generate(tid, approval(kind="technical_docx", task_id=task["id"]))
    artifact_id = repo.get_evidence(tid, source)["artifact_id"]
    if mutation == "source_bytes":
        repo.object_path(tid, artifact_id).write_bytes(b"changed")
    elif mutation == "source_missing":
        repo.object_path(tid, artifact_id).unlink()
    else:
        outputs.path(tid, output["id"]).unlink()
    with pytest.raises(ValueError, match="changed|missing|integrity"):
        release(submission_service(repo), tid, output)


def test_frozen_export_download_retains_approved_history_after_source_revision(workspace):
    repo, tid, _, outputs = workspace
    output = outputs.generate(tid, approval(kind="programme_xlsx", programme=programme()))
    service = submission_service(repo)
    package = release(service, tid, output)
    frozen = service.path(tid, package["id"]).read_bytes()
    seed(repo, tid, digest="b" * 64)
    assert service.path(tid, package["id"]).read_bytes() == frozen
    with pytest.raises(KeyError):
        service.path(repo.create_tender("Other")["id"], package["id"])


def test_output_publication_rejects_concurrent_record_change(workspace, monkeypatch):
    repo, tid, _, outputs = workspace
    import quantix.outputs as module

    original = module._extra_excel

    def render_then_change(*args):
        original(*args)
        seed(repo, tid, digest="b" * 64)

    monkeypatch.setattr(module, "_extra_excel", render_then_change)
    with pytest.raises(ValueError, match="changed"):
        outputs.generate(tid, approval(kind="programme_xlsx", programme=programme()))
    assert outputs.list(tid) == []
    assert list(outputs.directory.iterdir()) == []


def test_submission_routes_publish_models_and_reject_unreviewed_approval(workspace):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    repo, tid, _, outputs = workspace
    app = FastAPI()
    app.include_router(importlib.import_module("quantix.submission_routes").create_router(repo))
    client = TestClient(app)
    output = outputs.generate(tid, approval(kind="programme_xlsx", programme=programme()))
    preview = client.post(
        f"/api/tenders/{tid}/submissions/preview", json={"output_ids": [output["id"]]}
    )
    assert preview.status_code == 200
    response = client.post(
        f"/api/tenders/{tid}/submissions",
        json=approval(output_ids=[output["id"]], fingerprint=preview.json()["fingerprint"]),
    )
    assert response.status_code == 422
    assert "SubmissionApproval" in app.openapi()["components"]["schemas"]


def test_comparison_links_saved_rate_to_supplier_reply_with_commercial_state(workspace):
    from quantix.correspondence import QuoteService

    repo, tid, estimates, outputs = workspace
    mail = QuoteService(repo)
    quote = mail.create_draft(
        tid, {"to": ["supplier@example.com"], "subject": "Concrete", "body": "Quote concrete"}
    )
    reply = mail.register_reply(
        tid,
        quote["id"],
        {
            "sender": "supplier@example.com",
            "received_at": "2026-09-06T10:00:00+00:00",
            "text": "Concrete EGP 1200 per m3 excluding VAT, pump excluded.",
        },
    )
    item = estimates.view(tid)["items"][0]
    values = {
        "unit_rate": "1200",
        "currency": "EGP",
        "tax_basis": "excluding_vat",
        "vat_percent": "14",
        "provenance": {
            "basis": "observed",
            "observed_on": "2026-09-06",
            "source_ids": reply["source_ids"],
            "urls": [],
            "geography": "Cairo",
            "conditions": "Pump excluded",
        },
    }
    estimates.propose_rate(tid, item["id"], values)
    output = outputs.generate(tid, approval(kind="comparison_xlsx"))
    book = load_workbook(outputs.path(tid, output["id"]))
    rows = str(list(book["Rate comparison"].values))
    assert "supplier@example.com" in rows and "1200" in rows and "proposed" in rows
    assert "Pump excluded" in str(list(book["Rate provenance"].values))
    assert "unapproved" in " ".join(
        submission_service(repo).preview(tid, {"output_ids": [output["id"]]})["blocking_reasons"]
    )


def test_long_supplier_reply_is_preserved_across_continuation_rows(workspace):
    from quantix.correspondence import QuoteService

    repo, tid, _, outputs = workspace
    mail = QuoteService(repo)
    quote = mail.create_draft(
        tid, {"to": ["supplier@example.com"], "subject": "Concrete", "body": "Quote concrete"}
    )
    content = "Conditions and exclusions\n" * 1500 + "End of saved offer"
    mail.register_reply(
        tid,
        quote["id"],
        {
            "sender": "supplier@example.com",
            "received_at": "2026-09-06T10:00:00+00:00",
            "text": content,
        },
    )
    output = outputs.generate(tid, approval(kind="comparison_xlsx"))
    sheet = load_workbook(outputs.path(tid, output["id"]))["Supplier replies"]
    assert (
        "".join(str(sheet.cell(row, 5).value or "") for row in range(7, sheet.max_row + 1))
        == content
    )
