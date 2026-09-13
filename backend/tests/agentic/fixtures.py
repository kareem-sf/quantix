"""Reusable synthetic documents and controlled boundaries for acceptance drivers.

Every named colleague in these fixtures is generated test data, not a product
staff roster.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import Workbook

GENERATED_TEST_COLLEAGUE = "Generated test colleague: Amira Hassan (test data only)"


def pdf_bytes(texts: list[str | None]) -> bytes:
    """Small valid PDF with independent text objects and no third-party writer."""

    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    page_ids = []
    for text in texts:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        stream = (f"BT /F1 12 Tf 30 80 Td ({text}) Tj ET" if text else "").encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 100] /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    kids = " ".join(f"{n} 0 R" for n in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{n} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(data)


def write_pdf(path: Path, texts: list[str | None]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(pdf_bytes(texts))
    return path


def write_drawing(path: Path) -> Path:
    return write_pdf(path, ["Synthetic drawing SCALE 1:100", "Gridline A / Level 02"])


def write_docx(path: Path, paragraphs: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(GENERATED_TEST_COLLEAGUE)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)
    return path


def write_quotation(path: Path) -> Path:
    return write_docx(
        path,
        [
            "Quotation Q-TEST-001 for ready-mixed concrete.",
            "Item C20/25, 25 m3, unit rate recorded as a synthetic figure only.",
        ],
    )


def write_xlsx(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Civil"
    sheet.append(["Item", "Description", "Unit", "Quantity"])
    sheet.append(["C1", "Reinforced concrete slab (synthetic test row)", "m3", 28])
    book.save(path)
    book.close()
    return path


def write_xlsm(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Preliminaries"
    sheet.append(["Item", "Description", "Unit", "Quantity"])
    sheet.append(["P1", "Site establishment (synthetic test row)", "item", 1])
    book.save(path)
    book.close()
    return path


def write_source_revision_pair(directory: Path) -> tuple[Path, Path]:
    """Two synthetic revisions of the same relative specification."""

    first = write_pdf(directory / "rev-a" / "Specification.pdf", ["Concrete grade C30/37"])
    second = write_pdf(directory / "rev-b" / "Specification.pdf", ["Concrete grade C32/40"])
    return first, second


def write_package(directory: Path) -> dict[str, Path]:
    """A small mixed package used by isolated-home import cases."""

    directory.mkdir(parents=True, exist_ok=True)
    files = {
        "pdf": write_pdf(directory / "Specification.pdf", ["Reinforced concrete C30/37"]),
        "drawing": write_drawing(directory / "Drawings" / "GA-01.pdf"),
        "xlsx": write_xlsx(directory / "BOQ.xlsx"),
        "xlsm": write_xlsm(directory / "Preliminaries.xlsm"),
        "quotation": write_quotation(directory / "Quotes" / "Q-TEST-001.docx"),
    }
    return files


@dataclass
class RecordingProvider:
    """Controlled provider boundary that records calls and never contacts a network."""

    calls: list[dict] = field(default_factory=list)
    automatic: bool = False

    async def execute(self, *args, **kwargs):
        self.calls.append({"args": args, "kwargs": kwargs})
        if not self.automatic:
            raise AssertionError("This acceptance case must not start AI work.")
        raise AssertionError("A live provider is not part of this synthetic acceptance case.")


@dataclass
class FakeClock:
    instant: str = "2026-09-10T12:00:00.000+00:00"
    today: date = date(2026, 9, 10)


def seed_synthetic_office(repo, tender_id: str, *, model_id="synthetic-model"):
    """Admit a session-only synthetic account so Manager work can run without a network."""

    from quantix.ai_connections import AIConnectionService
    from quantix.ai_policy import AIPolicyService
    from quantix.ai_setup_store import SetupStore

    connections = AIConnectionService(repo)
    connection = connections.create(
        {
            "name": "Synthetic API",
            "provider_id": "openai",
            "protocol": "openai_chat",
            "base_url": "https://api.example.test/v1",
            "auth_type": "api_key",
            "billing": "metered",
            "credentials": {"api_key": "synthetic-key"},
            "session_only": True,
        }
    )
    connections.save_model(
        connection["id"],
        {
            "model_id": model_id,
            "display_name": "Synthetic model",
            "capabilities": {
                "tools": True,
                "structured_output": True,
                "web_search": True,
                "reasoning": ["high"],
                "max_output_tokens": 20000,
                "context_window": 200000,
            },
            "pricing": {
                "input_per_million": 1,
                "output_per_million": 2,
                "web_search_per_call": 0.001,
                "source": "synthetic",
                "as_of": "2026-09-10",
            },
        },
        source="provider",
    )
    connection = connections.get(connection["id"])
    store = SetupStore(repo)
    evidence = {
        "check": {"status": "passed", "model_id": model_id},
        "checked_revision": connection["revision"],
        "checked_component_version": "synthetic",
    }
    store.update(connection["id"], selected_model_id=model_id, **evidence)
    store.save_model_check(connection["id"], model_id, evidence)
    route = {
        "connection_id": connection["id"],
        "model_id": model_id,
        "reasoning": "high",
        "max_output_tokens": 2048,
        "web_search": False,
        "max_search_calls": 3,
    }
    AIPolicyService(repo).update(
        tender_id,
        {
            "allowed_connection_ids": [connection["id"]],
            "manager": route,
            "specialist": route,
            "run_budget_usd": 10,
            "tender_budget_usd": 100,
            "max_requests": 12,
            "engineer_confirmed": True,
            "rationale": "Synthetic approval",
        },
    )
    return connection, route


def prepare_busy_staff(
    repo,
    tender_id: str,
    *,
    tool_ids=("read_source",),
    max_depth=1,
    max_assignments=1,
    extra_work_orders=(),
):
    """Create generated staff with a queued assignment for lifecycle tests."""

    import json

    from quantix.ai_connections import AIConnectionService
    from quantix.ai_policy import AIPolicyService
    from quantix.db import new_id
    from quantix.manager_profile import ManagerProfileService
    from quantix.manager_runtime import ManagerRunProfiles
    from quantix.staff_assignments import StaffAssignmentService
    from quantix.staff_capabilities import get_capability
    from quantix.staff_models import ManagerCreationContext
    from quantix.staff_routing import StaffRoutingService
    from quantix.staff_routing_models import (
        ArtifactBasis,
        DelegationEnvelope,
        DelegationRouteOption,
        model_revision_fingerprint,
        route_option_id,
    )
    from quantix.staff_store import StaffStore

    connection, route = seed_synthetic_office(repo, tender_id)
    connections = AIConnectionService(repo)
    account = connections.get(connection["id"])
    model = next(item for item in connections.models(connection["id"]) if item["model_id"] == route["model_id"])
    body = b"synthetic reviewed source"
    digest = hashlib.sha256(body).hexdigest()
    (repo.objects / digest).write_bytes(body)
    artifact, _ = repo.register_artifact(
        tender_id,
        "Sources/spec.pdf",
        digest,
        len(body),
        {
            "kind": "pdf",
            "status": "extracted",
            "segments": [{"locator": "page:1", "text": body.decode()}],
        },
    )
    evidence = repo.artifact_evidence(tender_id, artifact["id"])[0]
    plan = repo.create_plan(
        tender_id,
        "Review",
        [{"title": "Review", "role": "Any role", "description": "Review", "source_ids": [evidence["id"]]}],
    )
    repo.approve_plan(tender_id, plan["id"], "Synthetic plan approval")
    manager = ManagerProfileService(repo)
    planning = repo.create_run(tender_id, "conversation", "Create colleague")
    context = ManagerCreationContext(tender_id, planning["id"], manager.get().version, plan["id"])
    staff = StaffStore(repo).create_generated(
        context,
        {
            "display_name": "Generated test colleague: Lina Haddad (test data only)",
            "role": "Package timing reviewer",
            "title": "Unregistered specialist",
            "specialisms": ["Package review"],
            "persona": "A generated test colleague.",
            "personality": {
                "description": "Evidence-led synthetic colleague.",
                "traits": ["careful"],
                "communication_style": "Plain language.",
                "problem_solving_style": "Use bounded checks.",
                "collaboration_style": "Share exact sources.",
                "uncertainty_handling": "Label gaps.",
                "initiative": "Suggest the next safe step.",
                "explanation_style": "Lead with the answer.",
                "language_preferences": ["English"],
                "working_habits": ["Keep citations"],
            },
            "responsibilities": ["Compare source evidence"],
            "objectives": ["List material gaps"],
            "methods": ["Read current sources"],
            "deliverables": ["A source-backed gap list"],
            "success_criteria": ["Every gap has a source"],
            "context_needs": ["Current Tender sources"],
            "requested_tool_ids": list(tool_ids),
            "creation_reason": "Created for this synthetic Tender task.",
        },
        {
            "brief": "Review the current Tender source package.",
            "goal": "Give the Manager a source-backed gap list.",
            "source_ids": [evidence["id"]],
            "expected_outputs": ["Gap list"],
            "completion_checks": ["Each gap has a source"],
        },
        "t003-create",
    )
    root = repo.create_run(tender_id, "manager", "Run delegated work")
    ManagerRunProfiles(repo).capture(tender_id, root["id"])
    option = DelegationRouteOption(
        id="0" * 64,
        route=route,
        connection_revision=account["revision"],
        model_revision=model_revision_fingerprint(model),
        model=model,
        account_name=account["name"],
        provider=account["provider_id"],
        data_destination=account["base_url"],
        billing=account["billing"],
        provider_managed_extras=False,
    )
    option = option.model_copy(update={"id": route_option_id(option)})
    envelope = DelegationEnvelope(
        version=1,
        purpose="Review current Tender sources.",
        source_scope="selected_sources",
        artifacts=[
            ArtifactBasis(
                artifact_id=artifact["id"],
                version=artifact["version"],
                content_hash=artifact["content_hash"],
            )
        ],
        tools=[get_capability(item) for item in tool_ids],
        allowed_draft_outputs=["findings"],
        route_options=[option],
        max_staff=1,
        max_assignments=max_assignments,
        max_depth=max_depth,
        max_concurrency=1,
        max_requests=12,
        max_search_calls=0,
        run_budget_usd=10,
        tender_budget_usd=100,
    )
    with repo.db.connect(write=True) as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS plan_review_approvals(id TEXT PRIMARY KEY,tender_id TEXT NOT NULL,plan_id TEXT NOT NULL,fingerprint TEXT NOT NULL,rationale TEXT NOT NULL,review_json TEXT NOT NULL,work_intents_json TEXT NOT NULL,created_at TEXT NOT NULL)"
        )
        conn.execute(
            "INSERT INTO plan_review_approvals VALUES(?,?,?,?,?,?,?,datetime('now'))",
            (
                new_id(),
                tender_id,
                plan["id"],
                "f" * 64,
                "Synthetic",
                json.dumps({"delegation": envelope.model_dump(mode="json")}),
                json.dumps(
                    [
                        {
                            "run_id": root["id"],
                            "tender_id": tender_id,
                            "plan_id": plan["id"],
                            "kind": "manager",
                            "task_id": None,
                        }
                    ]
                ),
            ),
        )
    routing = StaffRoutingService(repo)
    policy_revision = AIPolicyService(repo).get(tender_id)["revision"]
    routing.save_reviewed_grant(tender_id, plan["id"], "f" * 64, policy_revision, envelope)
    root_context = ManagerCreationContext(tender_id, root["id"], manager.get().version, plan["id"])
    binding = routing.bind(
        root_context, staff.staff.id, staff.work_order.id, envelope.route_options[0].id, "t003-bind"
    )
    assignment = StaffAssignmentService(repo).queue(root_context, binding.id, "t003-queue")
    store = StaffStore(repo)
    for index, order in enumerate(extra_work_orders):
        store.create_work_order(
            context, staff.staff.id, staff.staff.version, order, f"t008-extra-order-{index}"
        )
    return staff, assignment, root_context
