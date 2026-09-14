"""Reader-to-retrieval structure for PDF, Word and spreadsheets."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from quantix.documents import extract_document
from quantix.repository import Repository
from quantix.retrieval_passages import passages_for_evidence


class TinyModel:
    def token_count(self, text):
        return len(text) // 4


def _register(repo, tender_id, path: Path, extraction):
    body = path.read_bytes()
    artifact, _ = repo.register_artifact(
        tender_id,
        path.name,
        hashlib.sha256(body).hexdigest(),
        len(body),
        {
            "kind": extraction.kind,
            "status": extraction.status,
            "warnings": extraction.warnings,
            "metadata": extraction.metadata,
            "segments": [
                {
                    "locator": segment.locator,
                    "text": segment.text,
                    "page": segment.page,
                    "sheet": segment.sheet,
                    "cell_range": segment.cell_range,
                    "kind": segment.kind,
                    "metadata": segment.metadata,
                }
                for segment in extraction.segments
            ],
        },
    )
    return artifact


def _pdf_bytes(texts: list[str]) -> bytes:
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    page_ids = []
    for text in texts:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        stream = f"BT /F1 12 Tf 30 80 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 100] /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    kids = " ".join(f"{n} 0 R" for n in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{n} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(data)


def test_docx_headings_and_table_coordinates_are_observed(tmp_path):
    path = tmp_path / "spec.docx"
    document = Document()
    document.add_paragraph("Roof waterproofing", style="Heading 1")
    document.add_paragraph("Provide a four millimetre SBS membrane.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Unit"
    table.cell(1, 0).text = "Membrane"
    table.cell(1, 1).text = "m2"
    document.save(path)
    extraction = extract_document(path, path.name)
    heading = next(segment for segment in extraction.segments if segment.kind == "heading")
    assert heading.metadata["heading_level"] == 1
    assert heading.metadata["heading"] == "Roof waterproofing"
    body = next(segment for segment in extraction.segments if "four millimetre" in segment.text)
    assert body.metadata["section_heading"] == "Roof waterproofing"
    assert body.metadata["block_kind"] == "word_paragraph"
    cell = next(segment for segment in extraction.segments if segment.text.strip() == "m2")
    assert cell.metadata["table_row"] == 2
    assert cell.metadata["table_cell"] == 2
    assert cell.metadata["header_uncertain"] is True
    repo = Repository(tmp_path / "home")
    tender = repo.create_tender("Word")
    artifact = _register(repo, tender["id"], path, extraction)
    evidence = next(
        row
        for row in repo.artifact_evidence(tender["id"], artifact["id"])
        if "four millimetre" in row["text"]
    )
    passage = next(passages_for_evidence(evidence, model=TinyModel()))
    assert "Roof waterproofing" in passage.embed_text
    assert passage.excerpt.startswith("Provide a four")


def test_spreadsheet_links_rows_to_observed_headers_and_keeps_formula_warnings(tmp_path):
    path = tmp_path / "boq.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Civil"
    sheet["A1"] = "Item"
    sheet["B1"] = "Unit"
    sheet["C1"] = "Quantity"
    sheet["A2"] = "Blinding"
    sheet["B2"] = "m3"
    sheet["C2"] = 12.5
    sheet["D2"] = "=C2*100"
    sheet["A3"] = "=#REF!"
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Section 3"
    sheet.row_dimensions[2].hidden = True
    workbook.save(path)
    extraction = extract_document(path, path.name)
    header = next(
        segment
        for segment in extraction.segments
        if segment.metadata.get("row_role") == "observed_header"
    )
    assert header.metadata["header_uncertain"] is False
    data = next(segment for segment in extraction.segments if "Blinding" in segment.text)
    labels = {item["label"] for item in data.metadata["header_cells"]}
    assert {"Item", "Unit", "Quantity"} <= labels
    assert data.metadata["row_hidden"] is True
    formula_row = next(
        segment
        for segment in extraction.segments
        if any(cell.get("formula") for cell in segment.metadata["cells"])
    )
    assert any(cell.get("formula") == "=C2*100" for cell in formula_row.metadata["cells"])
    merged = next(segment for segment in extraction.segments if "Section 3" in segment.text)
    assert merged.metadata["merged_ranges"]
    passage = next(
        passages_for_evidence({"text": data.text, "metadata": data.metadata}, model=TinyModel())
    )
    assert "Item" in passage.embed_text
    assert "Blinding" in passage.excerpt


def test_pdf_pages_are_flagged_structure_uncertain_and_keep_page_locators(tmp_path):
    path = tmp_path / "note.pdf"
    path.write_bytes(
        _pdf_bytes(["Unless stated otherwise torch-applied.", "Provide the roof membrane."])
    )
    extraction = extract_document(path, path.name)
    assert extraction.segments
    assert all(
        segment.metadata.get("structure_uncertain") is True for segment in extraction.segments
    )
    assert [segment.locator for segment in extraction.segments] == ["page:1", "page:2"]
    repo = Repository(tmp_path / "home")
    tender = repo.create_tender("PDF")
    artifact = _register(repo, tender["id"], path, extraction)
    rows = repo.artifact_evidence(tender["id"], artifact["id"])
    first, second = rows[0], rows[1]
    passage = next(passages_for_evidence(second, model=TinyModel(), previous=first))
    assert passage.context_uncertain is True
    assert "Unless stated otherwise" in passage.embed_text
    assert second["text"][passage.start : passage.end] == passage.excerpt


def test_long_arabic_excerpt_offsets_remain_exact():
    text = "يجب أن توفر مضخة الخرسانة إنتاجية لا تقل عن خمسة وسبعين متراً مكعباً. " * 40
    passages = list(passages_for_evidence({"text": text}, model=TinyModel()))
    assert len(passages) > 1
    for passage in passages:
        assert text[passage.start : passage.end] == passage.excerpt
        assert "مضخة" in passage.excerpt or passage.start > 0
