"""Run a real isolated multilingual retrieval benchmark through Quantix."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import threading
import time
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

os.environ.update(
    {
        "OMP_NUM_THREADS": "1",
        "OPENBLAS_NUM_THREADS": "1",
        "MKL_NUM_THREADS": "1",
        "NUMEXPR_NUM_THREADS": "1",
    }
)

from docx import Document
from huggingface_hub import HfApi

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "backend"))

from quantix.intake import import_package
from quantix.repository import Repository
from quantix.semantic import (
    MODEL_FILES,
    MODEL_FINGERPRINT,
    MODEL_NAME,
    MODEL_REVISION,
    SemanticService,
)

MODEL_SOURCE = f"https://huggingface.co/{MODEL_NAME}/tree/{MODEL_REVISION}"

DOCUMENTS = {
    "01-pump-en.docx": (
        "Concrete placing plant",
        "The concrete pump shall provide a minimum output of 70 cubic metres per hour at the placing point.",
    ),
    "02-pump-ar.docx": (
        "معدات صب الخرسانة",
        "يجب أن توفر مضخة الخرسانة إنتاجية لا تقل عن 75 متراً مكعباً في الساعة عند نقطة الصب.",
    ),
    "03-waterproofing-en.docx": (
        "Roof waterproofing",
        "Provide a four millimetre SBS modified bituminous waterproofing membrane to the roof slab.",
    ),
    "04-reinforcement-ar.docx": (
        "حديد التسليح",
        "يجب أن تكون قضبان حديد التسليح من الصنف B500D طبقاً للجدول الإنشائي.",
    ),
    "05-curing-en.docx": (
        "Concrete curing",
        "Maintain moist curing of structural concrete for at least seven continuous days.",
    ),
    "06-commercial-ar.docx": (
        "شروط السعر والتوريد",
        "سعر توريد الطوب لا يشمل ضريبة القيمة المضافة ويشمل النقل إلى موقع المشروع في القاهرة.",
    ),
}

EXPECTED_TEXT = {
    "pump_en": DOCUMENTS["01-pump-en.docx"][1],
    "pump_ar": DOCUMENTS["02-pump-ar.docx"][1],
    "waterproof_en": DOCUMENTS["03-waterproofing-en.docx"][1],
    "reinforcement_ar": DOCUMENTS["04-reinforcement-ar.docx"][1],
    "curing_en": DOCUMENTS["05-curing-en.docx"][1],
    "commercial_ar": DOCUMENTS["06-commercial-ar.docx"][1],
}

QUERIES = (
    ("pump_en", "أي بند يشترط إنتاجية 70 متراً مكعباً في الساعة لمضخة الخرسانة؟"),
    ("pump_ar", "Which concrete pump requirement specifies 75 cubic metres per hour?"),
    ("waterproof_en", "ما بند غشاء عزل السطح المعدل SBS بسماكة أربعة ملليمترات؟"),
    ("reinforcement_ar", "Which reinforcement clause specifies steel grade B500D?"),
    ("curing_en", "ما البند الذي يشترط معالجة الخرسانة رطباً لمدة سبعة أيام متواصلة؟"),
    ("commercial_ar", "Which brick price clause excludes VAT and includes delivery to Cairo?"),
)


def _sha256(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def _write_docx(path: Path, heading: str, fact: str) -> None:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(fact)
    document.save(path)


def _write_package(directory: Path, documents: dict[str, tuple[str, str]]) -> None:
    directory.mkdir(parents=True, exist_ok=False)
    for name, (heading, fact) in documents.items():
        _write_docx(directory / name, heading, fact)


def _import(repo: Repository, tender_id: str, source: Path) -> dict:
    run = repo.create_run(
        tender_id, "import", "Synthetic multilingual benchmark import"
    )
    result = import_package(repo, tender_id, source, run["id"], threading.Event())
    repo.update_run(run["id"], status="completed", progress=100, result=result)
    return result


def _expected_evidence(repo: Repository, tender_id: str) -> dict[str, dict]:
    by_text: dict[str, dict] = {}
    for artifact in repo.list_artifacts(tender_id, current_only=True):
        for evidence in repo.artifact_evidence(tender_id, artifact["id"]):
            for key, text in EXPECTED_TEXT.items():
                if text in evidence["text"]:
                    by_text[key] = {
                        "source_id": evidence["id"],
                        "artifact_id": artifact["id"],
                        "artifact_name": artifact["name"],
                        "artifact_version": artifact["version"],
                        "content_hash": artifact["content_hash"],
                        "locator": evidence["locator"],
                    }
    missing = sorted(set(EXPECTED_TEXT) - set(by_text))
    if missing:
        raise RuntimeError(
            f"The actual import did not extract expected facts: {missing}"
        )
    return by_text


def _evaluate(service: SemanticService, tender_id: str, expected: dict[str, dict]):
    cases = []
    for key, query in QUERIES:
        started = time.perf_counter()
        hits = service.search(tender_id, query, limit=10)
        elapsed_ms = round((time.perf_counter() - started) * 1000, 2)
        identifiers = [item["id"] for item in hits]
        wanted = expected[key]["source_id"]
        rank = identifiers.index(wanted) + 1 if wanted in identifiers else None
        cases.append(
            {
                "case": key,
                "query": query,
                "expected": expected[key],
                "rank": rank,
                "top_source_ids": identifiers[:5],
                "top_scores": [round(float(item["score"]), 6) for item in hits[:5]],
                "elapsed_ms": elapsed_ms,
            }
        )
    count = len(cases)
    return {
        "cases": cases,
        "recall_at_1": sum(item["rank"] == 1 for item in cases) / count,
        "recall_at_3": sum(
            item["rank"] is not None and item["rank"] <= 3 for item in cases
        )
        / count,
        "recall_at_10": sum(item["rank"] is not None for item in cases) / count,
        "mean_reciprocal_rank": sum(
            0 if item["rank"] is None else 1 / item["rank"] for item in cases
        )
        / count,
    }


def _model_record(service: SemanticService) -> dict:
    info = HfApi(token=False).model_info(
        MODEL_NAME, revision=MODEL_REVISION, files_metadata=True
    )
    if info.sha != MODEL_REVISION:
        raise RuntimeError(
            f"The model registry resolved {MODEL_REVISION} to an unexpected commit {info.sha}."
        )
    remote = {}
    for sibling in info.siblings:
        if sibling.rfilename not in MODEL_FILES:
            continue
        lfs = sibling.lfs or {}
        expected_sha = (
            lfs.get("sha256") if isinstance(lfs, dict) else getattr(lfs, "sha256", None)
        )
        remote[sibling.rfilename] = {
            "blob_id": sibling.blob_id,
            "lfs_sha256": expected_sha,
            "remote_bytes": sibling.size,
        }
    files = {}
    for relative in MODEL_FILES:
        path = service.model_path / relative
        if not path.is_file():
            raise RuntimeError(
                f"The configured model file is missing after download: {relative}"
            )
        local_sha = _sha256(path)
        expected_sha = remote.get(relative, {}).get("lfs_sha256")
        if expected_sha and local_sha != expected_sha:
            raise RuntimeError(
                f"The downloaded model checksum differs from the registry: {relative}"
            )
        files[relative] = {
            "bytes": path.stat().st_size,
            "sha256": local_sha,
            "registry": remote.get(relative),
        }
    return {
        "name": MODEL_NAME,
        "revision": MODEL_REVISION,
        "model_fingerprint": MODEL_FINGERPRINT,
        "source": MODEL_SOURCE,
        "source_commit_verified": True,
        "files": files,
    }


def run(home: Path) -> dict:
    home = home.expanduser().resolve()
    home.mkdir(parents=True, exist_ok=True)
    input_root = home.parent / "retrieval-synthetic-input"
    input_root.mkdir(parents=True, exist_ok=True)
    package = input_root / uuid4().hex
    _write_package(package, DOCUMENTS)

    started = time.perf_counter()
    repo = Repository(home)
    tender = repo.create_tender(
        f"Synthetic multilingual retrieval {datetime.now(UTC).isoformat()}"
    )
    imported = _import(repo, tender["id"], package)
    expected = _expected_evidence(repo, tender["id"])
    service = SemanticService(repo)
    before = service.status(tender["id"])
    index_started = time.perf_counter()
    indexed = service.index(tender["id"])
    index_seconds = round(time.perf_counter() - index_started, 3)
    evaluation = _evaluate(service, tender["id"], expected)

    old_source_id = expected["pump_en"]["source_id"]
    revised_fact = (
        "The concrete pump shall provide a minimum output of 82 cubic metres per hour "
        "at the placing point."
    )
    revised = dict(DOCUMENTS)
    revised["01-pump-en.docx"] = ("Concrete placing plant", revised_fact)
    _write_docx(package / "01-pump-en.docx", *revised["01-pump-en.docx"])
    revision_import = _import(repo, tender["id"], package)
    stale_status = service.status(tender["id"])
    revision_index_started = time.perf_counter()
    revised_index = service.index(tender["id"])
    revision_index_seconds = round(time.perf_counter() - revision_index_started, 3)
    revised_artifact = next(
        item
        for item in repo.list_artifacts(tender["id"], current_only=True)
        if item["name"] == "01-pump-en.docx"
    )
    revised_evidence = next(
        item
        for item in repo.artifact_evidence(tender["id"], revised_artifact["id"])
        if revised_fact in item["text"]
    )
    revision_hits = service.search(
        tender["id"],
        "أي بند يشترط إنتاجية 82 متراً مكعباً في الساعة لمضخة الخرسانة؟",
        limit=10,
    )
    revision_ids = [item["id"] for item in revision_hits]

    output = home / "benchmarks" / f"multilingual-{uuid4().hex}.json"
    result = {
        "schema_version": 1,
        "run_at": datetime.now(UTC).isoformat(),
        "isolated_home": str(home),
        "synthetic_input": str(package),
        "tender_id": tender["id"],
        "model": _model_record(service),
        "import": imported,
        "status_before_index": before,
        "indexed": indexed,
        "index_seconds": index_seconds,
        "evaluation": evaluation,
        "revision": {
            "import": revision_import,
            "status_before_reindex": stale_status,
            "indexed": revised_index,
            "index_seconds": revision_index_seconds,
            "old_source_id": old_source_id,
            "new_source_id": revised_evidence["id"],
            "new_locator": revised_evidence["locator"],
            "new_artifact_version": revised_artifact["version"],
            "new_content_hash": revised_artifact["content_hash"],
            "old_source_excluded": old_source_id not in revision_ids,
            "new_source_rank": revision_ids.index(revised_evidence["id"]) + 1
            if revised_evidence["id"] in revision_ids
            else None,
            "top_source_ids": revision_ids[:5],
        },
        "total_seconds": round(time.perf_counter() - started, 3),
        "result_path": str(output),
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    sys.stdout.buffer.write(
        (json.dumps(result, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--home",
        type=Path,
        default=Path.home()
        / ".quantix"
        / "cache"
        / "development"
        / "2026-09-12-full-agentic-office"
        / "retrieval",
    )
    args = parser.parse_args()
    run(args.home)


if __name__ == "__main__":
    main()
